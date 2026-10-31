"""Formula results stored in the file, so a viewer that does not calculate
still shows them."""

import os
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from prometheus_core.xlsx_cache import fill_sheet_xml, store_formula_results

try:
    import openpyxl
except Exception:  # pragma: no cover - optional dependency
    openpyxl = None


class FillXmlTests(unittest.TestCase):
    def test_a_number_goes_into_the_empty_slot(self):
        xml, n = fill_sheet_xml('<c r="B1" s="2"><f>A1*2</f><v /></c>', {"B1": 10.5})
        self.assertEqual(xml, '<c r="B1" s="2"><f>A1*2</f><v>10.5</v></c>')
        self.assertEqual(n, 1)

    def test_an_empty_result_is_stored_as_an_empty_string(self):
        xml, _ = fill_sheet_xml('<c r="C1"><f>IF(A1=0,"",A1)</f><v /></c>', {"C1": ""})
        self.assertEqual(xml, '<c r="C1" t="str"><f>IF(A1=0,"",A1)</f><v></v></c>')

    def test_a_cell_not_asked_for_is_left_alone(self):
        src = '<c r="D1"><f>A1</f><v /></c>'
        self.assertEqual(fill_sheet_xml(src, {"B1": 1}), (src, 0))

    def test_a_value_already_stored_is_never_overwritten(self):
        src = '<c r="B1"><f>A1*2</f><v>99</v></c>'
        self.assertEqual(fill_sheet_xml(src, {"B1": 1}), (src, 0))

    def test_text_is_escaped(self):
        xml, _ = fill_sheet_xml('<c r="A1"><f>X</f><v /></c>', {"A1": "a<b & c"})
        self.assertIn("<v>a&lt;b &amp; c</v>", xml)

    def test_a_cell_name_is_matched_whole_not_as_a_prefix(self):
        src = '<c r="B10"><f>A1</f><v /></c>'
        self.assertEqual(fill_sheet_xml(src, {"B1": 5}), (src, 0))


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class WorkbookTests(unittest.TestCase):
    def build(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws["A1"] = 5
        ws["B1"] = "=A1*2"
        ws["C1"] = '=IF(A1=0,"",A1)'
        other = wb.create_sheet("Daily detail")
        other["A1"] = "=1+1"
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(path)
        self.addCleanup(os.remove, path)
        return path

    def test_stored_results_read_back_without_calculating(self):
        path = self.build()
        n = store_formula_results(path, {"Summary": {"B1": 10, "C1": 5},
                                         "Daily detail": {"A1": 2}})
        self.assertEqual(n, 3)
        values = openpyxl.load_workbook(path, data_only=True)
        self.assertEqual(values["Summary"]["B1"].value, 10)
        self.assertEqual(values["Summary"]["C1"].value, 5)
        self.assertEqual(values["Daily detail"]["A1"].value, 2)

    def test_the_formulas_themselves_are_kept(self):
        path = self.build()
        store_formula_results(path, {"Summary": {"B1": 10}})
        formulas = openpyxl.load_workbook(path)
        self.assertEqual(formulas["Summary"]["B1"].value, "=A1*2")

    def test_the_rest_of_the_package_is_carried_over_untouched(self):
        path = self.build()
        before = {i.filename for i in zipfile.ZipFile(path).infolist()}
        store_formula_results(path, {"Summary": {"B1": 10}})
        after = {i.filename for i in zipfile.ZipFile(path).infolist()}
        self.assertEqual(before, after)

    def test_an_unknown_sheet_is_ignored(self):
        path = self.build()
        self.assertEqual(store_formula_results(path, {"Nope": {"A1": 1}}), 0)


if __name__ == "__main__":
    unittest.main()
