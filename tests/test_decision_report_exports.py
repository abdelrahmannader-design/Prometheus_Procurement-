"""The two exports of the decision report: one model, two faithful renderings."""

import datetime as dt
import os
import re
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from prometheus_core import timing_report as tr
from prometheus_core.report_html import render

try:
    import openpyxl
    from prometheus_core.report_excel import SHEETS, build_workbook
except Exception:  # pragma: no cover - optional dependency
    openpyxl = None

F = 0.3937


def model():
    hist = tr.Series([(dt.date(2026, 3, 1) + dt.timedelta(days=i), 450.0 + i) for i in range(60)])
    contracts = [
        {"commodity": "CORN-BRZ", "status": "Closed", "qty_mt": 5000, "premium_cents": 190.0,
         "pricing_date": "2026-03-31", "cif_usd_mt": (470 + 190) * F, "priced": True},
        {"commodity": "SFM", "status": "Closed", "qty_mt": 2000, "cif_usd_mt": 282.0, "priced": True},
        {"commodity": "CORN-BRZ", "status": "Closed", "qty_mt": 3000, "premium_cents": 180.0,
         "cif_usd_mt": 255.0, "priced": True},
    ]
    rows = []
    for i, c in enumerate(contracts):
        base = tr.base_commodity(c["commodity"])
        factor = tr.DEFAULT_FACTORS.get(base)
        fx = 50.0
        landed = c["cif_usd_mt"] * fx + 900
        local = landed + (300 if i != 2 else -400)
        timing = tr.contract_timing(c, hist if factor else tr.Series([]), factor, cost_fx=fx)
        local_a = tr.assess_local(local - landed, local)
        rows.append({
            "id": f"C{i}", "contract": ["Brz </script> cargo", "SFM lot", "No date"][i],
            "supplier": "S", "commodity": c["commodity"], "base": base, "origin": "BRAZIL",
            "status": "Closed", "pricing_status": "PRICED", "qty_mt": c["qty_mt"],
            "signing_date": "", "shipment_date": "", "delivery_date": "2026-04-10",
            "pricing_date": c.get("pricing_date", ""), "form4_date": "",
            "cif_usd_mt": c["cif_usd_mt"], "fx": fx, "fx_source": "delivery",
            "contract_value_usd": c["qty_mt"] * c["cif_usd_mt"],
            "contract_value_egp": c["qty_mt"] * c["cif_usd_mt"] * fx,
            "discharge_egp_mt": 300.0, "clearance_egp_mt": 200.0, "freight_egp_mt": 400.0,
            "landed_cost_egp_mt": landed, "local_egp_mt": local, "local_key": "CORN-BRZ",
            "local_date": "2026-04-09", "saving_local_egp_mt": local - landed,
            "saving_local_egp": (local - landed) * c["qty_mt"], "saving_kind": "Realized",
            "local_assessment": local_a, "premium": c.get("premium_cents"),
            "futures_month": "JUL", "cbot_factor": factor, "cbot_at_pricing": 480.0,
            "cbot_at_delivery": 489.0, "cbot_latest": 509.0, "cbot_latest_date": "2026-04-29",
            "daily_consumption_mt": 500.0, "stock_before_arrival_mt": 1000.0,
            "coverage": tr.arrival_coverage(1000.0, 500.0, c["qty_mt"]),
            "timing": timing, "final_assessment": tr.final_assessment(local_a, timing["assessment"]),
            "missing": ["contract signing date", "shipment date", "shipment lead time"],
        })
    return {"title": "Procurement Decision & Contract Timing Report", "generated": "now",
            "as_of": "2026-04-29", "app": "test", "commodity": "ALL", "params": dict(tr.DEFAULTS),
            "contracts": rows, "summary": tr.executive_summary(rows),
            "stock": [{"base": "CORN", "daily_rate": 500.0, "stock_mt": 1000.0, "cover_days": 2.0,
                       "inbound_mt": 0.0, "cover_with_inbound_days": 2.0, "shortage_mt": 0.0}],
            "market": {"CORN": [[d.isoformat(), v] for d, v in zip(hist.dates, hist.values)]},
            "fx": [], "gaps": [{"id": "C2", "contract": "No date", "missing": "pricing date"}],
            "version": tr.TIMING_REPORT_VERSION}


class HtmlTests(unittest.TestCase):
    def test_it_fetches_nothing_from_anywhere(self):
        h = render(model())
        self.assertFalse(re.search(r'(?:src|href)=["\'](?:https?:)?//', h))
        self.assertNotIn("<link", h)
        self.assertNotIn("<script src", h)

    def test_the_data_cannot_close_its_own_script_block(self):
        h = render(model())
        block = h.split('id="report-data">', 1)[1].split("</script>", 1)[0]
        self.assertIn("Brz <\\/script> cargo", block)

    def test_every_contract_is_in_the_embedded_data(self):
        h = render(model())
        for cid in ("C0", "C1", "C2"):
            self.assertIn(f'"id":"{cid}"', h)

    def test_it_says_hindsight_is_hindsight(self):
        self.assertIn("hindsight benchmarks", render(model()))


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class ExcelTests(unittest.TestCase):
    def build(self):
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        self.addCleanup(os.remove, path)
        stored = build_workbook(model(), path)
        return path, stored

    def test_the_eight_sheets_in_order(self):
        path, _ = self.build()
        self.assertEqual(openpyxl.load_workbook(path).sheetnames, list(SHEETS))

    def test_every_formula_carries_a_stored_result(self):
        path, stored = self.build()
        with zipfile.ZipFile(path) as z:
            empty = sum(len(re.findall(r"</f><v\s*/></c>", z.read(n).decode()))
                        for n in z.namelist() if n.startswith("xl/worksheets/sheet"))
        self.assertGreater(stored, 100)
        self.assertEqual(empty, 0)

    def test_the_timing_opportunity_is_a_formula_with_the_engine_value(self):
        path, _ = self.build()
        f = openpyxl.load_workbook(path)["CBOT & Timing"]
        v = openpyxl.load_workbook(path, data_only=True)["CBOT & Timing"]
        self.assertTrue(str(f["S5"].value).startswith("="))
        self.assertAlmostEqual(v["S5"].value, model()["contracts"][0]["timing"]["opportunity_usd"], places=6)

    def test_missing_timing_is_written_as_insufficient_not_zero(self):
        path, _ = self.build()
        v = openpyxl.load_workbook(path, data_only=True)["CBOT & Timing"]
        self.assertEqual(v["AB7"].value, tr.INSUFFICIENT)
        self.assertIn(v["S7"].value, (None, ""))

    def test_the_simulator_inputs_are_editable_cells(self):
        path, _ = self.build()
        x = openpyxl.load_workbook(path)["Scenario Simulator"]
        for ref in ("C4", "C5", "C6", "C7", "C9", "C10", "C11", "C12"):
            self.assertIsInstance(x[ref].value, (int, float), ref)

    def test_savings_and_timing_are_separate_cards(self):
        path, _ = self.build()
        e = openpyxl.load_workbook(path)["Executive Summary"]
        labels = [c.value for row in e.iter_rows(min_row=5, max_row=13) for c in row if isinstance(c.value, str)]
        self.assertIn("Saving vs local — realized, EGP", labels)
        self.assertIn("Timing opportunity (hindsight), USD", labels)
        self.assertFalse(any("total saving" in l.lower() for l in labels))


if __name__ == "__main__":
    unittest.main()
