"""Excel workbook for the Procurement Decision & Contract Timing Report.

Takes the report model built by the app (see ``_tr_build_model``) and
writes eight sheets. Figures a reader may want to audit or change are live
formulas over the Contract Register and the Scenario Simulator, so moving a
threshold or an assumption re-runs the dependent results. Each formula's
result is also stored in the file (xlsx_cache), so the workbook reads
correctly in viewers that do not calculate.
"""

from __future__ import annotations

import datetime as dt

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart
from openpyxl.chart import Series as ChartSeries
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import timing_report as tr
from .xlsx_cache import store_formula_results

REPORT_EXCEL_VERSION = "1.0"

NAVY, BLUE, GREY = "1F3A5F", "1A6EBD", "8A8A8A"
GREEN_BG, AMBER_BG, RED_BG, GREY_BG = "E3F4E1", "FFF4D6", "FBE3E1", "EFEFEF"
MONEY0, MONEY2, PCT, INT = "#,##0", "#,##0.00", "0.00%", "#,##0"

SHEETS = ("Executive Summary", "Contract Register", "Savings vs Local",
          "Consumption Analysis", "CBOT & Timing", "Missed Opportunities",
          "Scenario Simulator", "Data & Methodology")


def _q(title):
    return "'" + title.replace("'", "''") + "'"


def _blank(v):
    return "" if v is None else v


class _Book:
    """A workbook that remembers the value of every formula it writes."""

    def __init__(self):
        self.wb = Workbook()
        self.cache = {}
        first = self.wb.active
        first.title = SHEETS[0]
        for title in SHEETS[1:]:
            self.wb.create_sheet(title)

    def __getitem__(self, title):
        return self.wb[title]

    def put(self, ws, row, col, value, fmt=None, bold=False, color=None, fill=None,
            wrap=False, italic=False, size=None):
        cell = ws.cell(row, col, value)
        if fmt:
            cell.number_format = fmt
        if bold or color or italic or size:
            cell.font = Font(bold=bold, color=color, italic=italic, size=size)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if wrap:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        return cell

    def formula(self, ws, row, col, formula, value, fmt=None, bold=False):
        cell = self.put(ws, row, col, formula, fmt, bold)
        self.cache.setdefault(ws.title, {})[cell.coordinate] = "" if value is None else value
        return cell

    def header(self, ws, row, heads, start=1, height=34):
        for i, (label, width) in enumerate(heads):
            c = ws.cell(row, start + i, label)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            ws.column_dimensions[get_column_letter(start + i)].width = width
        ws.row_dimensions[row].height = height

    def title(self, ws, text, sub=None):
        self.put(ws, 1, 1, text, bold=True, size=15, color=NAVY)
        if sub:
            self.put(ws, 2, 1, sub, italic=True, color=GREY)

    def verdict_colours(self, ws, rng):
        first = rng.split(":")[0]
        for word, colour in (("Favorable", GREEN_BG), ("Acceptable", AMBER_BG),
                             ("Unfavorable", RED_BG), ("Insufficient Data", GREY_BG),
                             ("Needed", GREEN_BG), ("Early", AMBER_BG)):
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'ISNUMBER(SEARCH("{word}",{first}))'],
                fill=PatternFill("solid", fgColor=colour)))


def _weighted(events, key):
    good = [e for e in events if e.get("status") == "assessed"]
    q = sum(e["qty_mt"] for e in good)
    return (sum(e[key] * e["qty_mt"] for e in good) / q) if q else None


def build_workbook(model, path):
    """Write the report to ``path``; returns the number of stored results."""
    b = _Book()
    rows = model["contracts"]
    n = len(rows)
    p = model["params"]
    first, last = 5, 4 + n          # data rows on every per-contract sheet
    R, S, C, T, M, X, Dm, E = (b[s] for s in ("Contract Register", "Savings vs Local",
                                               "Consumption Analysis", "CBOT & Timing",
                                               "Missed Opportunities", "Scenario Simulator",
                                               "Data & Methodology", "Executive Summary"))
    reg, sav, con, tim, sim = (_q(s) for s in ("Contract Register", "Savings vs Local",
                                               "Consumption Analysis", "CBOT & Timing",
                                               "Scenario Simulator"))
    ex = _q("Executive Summary")

    # ── Scenario Simulator inputs first: other sheets read its thresholds ──
    b.title(X, "Scenario Simulator",
            "Change the yellow cells; every result below and the assessments on the other "
            "sheets recalculate. Base case = all shifts at 0.")
    inputs = [
        (4, "CBOT shift (¢/bu for grains, $/short ton for SBM) — premium contracts only", 0, "0.00"),
        (5, "Premium shift (same units as the CBOT)", 0, "0.00"),
        (6, "FX change, %", 0, "0.00"),
        (7, "Local price change, %", 0, "0.00"),
        (9, "Local result: ± band counted as Acceptable, % of local", p["local_band_pct"], "0.00"),
        (10, "Timing: opportunity below this % is Favorable", p["small_opportunity_pct"], "0.00"),
        (11, "Stock: cover below this many days was tight", p["safety_days"], "0"),
        (12, "Stock: cover above this many days arrived early", p["ample_days"], "0"),
    ]
    b.put(X, 3, 2, "Assumption", bold=True)
    b.put(X, 3, 3, "Value", bold=True)
    for r, label, value, fmt in inputs:
        b.put(X, r, 2, label)
        b.put(X, r, 3, value, fmt, fill="FFF2A8")
    b.put(X, 8, 2, "Thresholds", bold=True)
    b.put(X, 13, 2, f"Timing window: {p['window_before_days']} days before to "
                    f"{p['window_after_days']} days after each pricing day — set in the app "
                    f"and re-export to change (it needs the CBOT history).", italic=True, color=GREY)
    X.column_dimensions["B"].width = 58
    band, small, safety, ample = (f"{sim}!$C$9", f"{sim}!$C$10", f"{sim}!$C$11", f"{sim}!$C$12")

    # ── Contract Register ─────────────────────────────────────────────────
    b.title(R, "Contract Register", "Every contract as recorded in the app. Blue headers "
            "are inputs from the JSON; the two contract-value columns are formulas.")
    reg_cols = [
        ("id", "ID", 8), ("contract", "Contract", 24), ("supplier", "Supplier", 16),
        ("commodity", "Commodity", 11), ("base", "Base", 8), ("origin", "Origin", 12),
        ("status", "Status", 8), ("pricing_status", "Pricing status", 11),
        ("qty_mt", "Qty MT", 10), ("signing_date", "Signing date", 12),
        ("pricing_date", "Pricing date", 12), ("form4_date", "Form 4 date", 12),
        ("shipment_date", "Shipment date", 12), ("delivery_date", "Arrival date", 12),
        ("cif_usd_mt", "CIF USD/MT", 10), ("fx", "FX used", 9), ("fx_source", "FX source", 11),
        ("value_usd", "Contract value USD", 14), ("value_egp", "Contract value EGP", 16),
        ("premium", "Premium", 9), ("futures_month", "Futures month", 10),
        ("cbot_factor", "CBOT factor", 9), ("discharge_egp_mt", "Intake / discharge EGP/MT", 12),
        ("clearance_egp_mt", "Clearance EGP/MT", 11), ("freight_egp_mt", "Freight EGP/MT", 11),
        ("local_egp_mt", "Local EGP/MT (delivered)", 13), ("local_key", "Local price of", 11),
        ("local_date", "Local price date", 12), ("daily", "Daily consumption MT", 11),
        ("stock_before", "Stock before arrival MT", 13),
    ]
    rc = {key: get_column_letter(i + 1) for i, (key, _, _) in enumerate(reg_cols)}
    b.header(R, 4, [(h, w) for _, h, w in reg_cols])
    for i, row in enumerate(rows):
        r = first + i
        vals = dict(row, daily=row.get("daily_consumption_mt"),
                    stock_before=row.get("stock_before_arrival_mt"),
                    signing_date="not recorded", shipment_date="not recorded")
        for j, (key, _, _) in enumerate(reg_cols):
            col = j + 1
            if key == "value_usd":
                b.formula(R, r, col, f'=IF(OR({rc["qty_mt"]}{r}="",{rc["cif_usd_mt"]}{r}=""),"",'
                                     f'{rc["qty_mt"]}{r}*{rc["cif_usd_mt"]}{r})',
                          row.get("contract_value_usd"), MONEY0)
            elif key == "value_egp":
                b.formula(R, r, col, f'=IF(OR({rc["value_usd"]}{r}="",{rc["fx"]}{r}=""),"",'
                                     f'{rc["value_usd"]}{r}*{rc["fx"]}{r})',
                          row.get("contract_value_egp"), MONEY0)
            else:
                v = vals.get(key)
                fmt = (MONEY2 if key in ("cif_usd_mt", "discharge_egp_mt", "clearance_egp_mt",
                                         "freight_egp_mt", "local_egp_mt", "premium")
                       else "0.0000" if key in ("fx", "cbot_factor")
                       else MONEY0 if key in ("qty_mt", "daily", "stock_before") else None)
                b.put(R, r, col, _blank(v), fmt)
    tot = last + 1
    b.put(R, tot, 1, "TOTAL", bold=True)
    for key, value in (("qty_mt", model["summary"]["quantity_mt"]),
                       ("value_usd", model["summary"]["contract_value_usd"]),
                       ("value_egp", model["summary"]["contract_value_egp"])):
        col = rc[key]
        b.formula(R, tot, reg_cols.index(next(x for x in reg_cols if x[0] == key)) + 1,
                  f"=SUM({col}{first}:{col}{last})", value, MONEY0, bold=True)
    R.freeze_panes = "C5"
    R.auto_filter.ref = f"A4:{get_column_letter(len(reg_cols))}{last}"
    RG = lambda key, r: f"{reg}!{rc[key]}{r}"

    # ── Savings vs Local ──────────────────────────────────────────────────
    b.title(S, "Savings vs Local", "Landed cost = CIF × FX + intake/discharge + clearance + "
            "freight. Saving = delivered local price − landed cost. Positive = the import "
            "was cheaper. Realized = closed & priced; Estimated = open, marked to today.")
    b.header(S, 4, [("ID", 8), ("Contract", 24), ("Base", 8), ("Qty MT", 10), ("CIF USD/MT", 10),
                    ("FX used", 9), ("Intake / discharge", 11), ("Clearance", 10), ("Freight", 10),
                    ("Landed cost EGP/MT", 13), ("Local EGP/MT", 12), ("Saving EGP/MT", 12),
                    ("Saving EGP", 15), ("Saving % of local", 10), ("Kind", 14),
                    ("Local result", 14)])
    for i, row in enumerate(rows):
        r = first + i
        b.put(S, r, 1, row["id"]); b.put(S, r, 2, row["contract"]); b.put(S, r, 3, row["base"])
        refs = {4: "qty_mt", 5: "cif_usd_mt", 6: "fx", 7: "discharge_egp_mt",
                8: "clearance_egp_mt", 9: "freight_egp_mt"}
        for col, key in refs.items():
            v = row.get(key)
            b.formula(S, r, col, f'=IF({RG(key, r)}="","",{RG(key, r)})', v,
                      MONEY0 if key == "qty_mt" else ("0.0000" if key == "fx" else MONEY2))
        b.formula(S, r, 10, f'=IF(OR(E{r}="",F{r}=""),"",E{r}*F{r}+N(G{r})+N(H{r})+N(I{r}))',
                  row.get("landed_cost_egp_mt"), MONEY2)
        b.formula(S, r, 11, f'=IF({RG("local_egp_mt", r)}="","",{RG("local_egp_mt", r)})',
                  row.get("local_egp_mt"), MONEY2)
        sav_mt = row.get("saving_local_egp_mt")
        b.formula(S, r, 12, f'=IF(OR(J{r}="",K{r}=""),"",K{r}-J{r})', sav_mt, MONEY2)
        b.formula(S, r, 13, f'=IF(OR(L{r}="",D{r}=""),"",L{r}*D{r})', row.get("saving_local_egp"), MONEY0)
        local = row.get("local_egp_mt")
        b.formula(S, r, 14, f'=IF(OR(L{r}="",K{r}="",K{r}=0),"",L{r}/K{r})',
                  (sav_mt / local) if (sav_mt is not None and local) else None, PCT)
        b.put(S, r, 15, row.get("saving_kind"))
        b.formula(S, r, 16, f'=IF(N{r}="","Insufficient Data",IF(N{r}*100>={band},"Favorable",'
                            f'IF(N{r}*100<=-{band},"Unfavorable","Acceptable")))',
                  row.get("local_assessment"))
    b.put(S, tot, 1, "TOTAL", bold=True)
    b.formula(S, tot, 13, f"=SUM(M{first}:M{last})",
              sum(r["saving_local_egp"] for r in rows if r.get("saving_local_egp") is not None),
              MONEY0, bold=True)
    for k, (label, kind) in enumerate((("of which Realized", "Realized"),
                                       ("of which Estimated (open)", "Estimated"),
                                       ("Not in Savings Tracker", "Not in Savings Tracker"))):
        rr = tot + 1 + k
        b.put(S, rr, 12, label, italic=True)
        b.formula(S, rr, 13, f'=SUMIFS(M{first}:M{last},O{first}:O{last},"{kind}")',
                  sum(r["saving_local_egp"] for r in rows
                      if r.get("saving_kind") == kind and r.get("saving_local_egp") is not None),
                  MONEY0)
    b.verdict_colours(S, f"P{first}:P{last}")
    S.freeze_panes = "C5"
    S.auto_filter.ref = f"A4:P{last}"

    # ── Consumption Analysis ──────────────────────────────────────────────
    b.title(C, "Consumption Analysis", "Stock from the app's FIFO walk at the configured "
            "daily rate, every day from each delivery. Cover = stock ÷ daily consumption.")
    b.put(C, 3, 1, "STOCK TODAY BY COMMODITY", bold=True)
    b.header(C, 4, [("Base", 8), ("Daily consumption MT", 12), ("Stock now MT", 12),
                    ("Cover days", 10), ("Inbound MT", 12), ("Cover incl. inbound, days", 13),
                    ("Projected shortage MT", 13)])
    stock_rows = model.get("stock") or []
    for i, s_row in enumerate(stock_rows):
        r = 5 + i
        b.put(C, r, 1, s_row["base"])
        b.put(C, r, 2, _blank(s_row.get("daily_rate")), MONEY0)
        b.put(C, r, 3, s_row.get("stock_mt"), MONEY0)
        b.formula(C, r, 4, f'=IF(OR(B{r}="",B{r}=0),"",C{r}/B{r})', s_row.get("cover_days"), "0.0")
        b.put(C, r, 5, s_row.get("inbound_mt"), MONEY0)
        b.formula(C, r, 6, f'=IF(OR(B{r}="",B{r}=0),"",(C{r}+E{r})/B{r})',
                  s_row.get("cover_with_inbound_days"), "0.0")
        b.put(C, r, 7, s_row.get("shortage_mt"), MONEY0)
    note_r = 5 + len(stock_rows)
    b.put(C, note_r, 1, "Projected shortage = consumption at the daily rate on days when no "
                        "recorded cargo had stock left (opening stock not on file, or a rate "
                        "above what was bought).", italic=True, color=GREY)
    top = max(note_r + 2, 12)
    b.put(C, top - 1, 1, "AT EACH ARRIVAL", bold=True)
    b.header(C, top, [("ID", 8), ("Contract", 24), ("Base", 8), ("Arrival date", 12),
                      ("Qty MT", 10), ("Daily consumption MT", 11), ("Days this cargo covers", 11),
                      ("Stock before arrival MT", 12), ("Cover before arrival, days", 12),
                      ("Arrival verdict", 26), ("Max delay without shortage, days", 13),
                      ("Explanation", 70)])
    c_first = top + 1
    for i, row in enumerate(rows):
        r = c_first + i
        rr = first + i
        cov = row.get("coverage") or {}
        b.put(C, r, 1, row["id"]); b.put(C, r, 2, row["contract"]); b.put(C, r, 3, row["base"])
        b.put(C, r, 4, row.get("delivery_date"))
        b.formula(C, r, 5, f'=IF({RG("qty_mt", rr)}="","",{RG("qty_mt", rr)})', row.get("qty_mt"), MONEY0)
        b.formula(C, r, 6, f'=IF({RG("daily", rr)}="","",{RG("daily", rr)})',
                  row.get("daily_consumption_mt"), MONEY0)
        b.formula(C, r, 7, f'=IF(OR(F{r}="",F{r}=0,E{r}=""),"",E{r}/F{r})', cov.get("cargo_days"), "0.0")
        b.formula(C, r, 8, f'=IF({RG("stock_before", rr)}="","",{RG("stock_before", rr)})',
                  row.get("stock_before_arrival_mt"), MONEY0)
        b.formula(C, r, 9, f'=IF(OR(F{r}="",F{r}=0,H{r}=""),"",H{r}/F{r})', cov.get("cover_days"), "0.0")
        b.formula(C, r, 10, f'=IF(I{r}="","Insufficient Data",IF(I{r}<=0.5,"Needed — yard empty",'
                            f'IF(I{r}<{safety},"Needed — prevented shortage",IF(I{r}<={ample},'
                            f'"Timely","Early — stock was ample"))))', cov.get("verdict"))
        b.formula(C, r, 11, f'=IF(I{r}="","",MAX(I{r}-{safety},0))', cov.get("max_safe_delay_days"), "0")
        b.put(C, r, 12, cov.get("explain") or "", wrap=False)
    c_last = c_first + n - 1
    b.verdict_colours(C, f"J{c_first}:J{c_last}")
    C.freeze_panes = C.cell(c_first, 3)
    C.auto_filter.ref = f"A{top}:L{c_last}"

    # ── CBOT & Timing ─────────────────────────────────────────────────────
    b.title(T, "CBOT & Timing — when the CBOT was fixed",
            "Hindsight benchmark: the lowest CBOT in the window around each pricing day, "
            "moved onto the contract's own priced CBOT. It uses prices from after the "
            "decision, so it is not a claim that the best day was knowable then; the "
            "20-day columns use only prices on or before the pricing day.")
    b.header(T, 4, [("ID", 8), ("Contract", 22), ("Base", 7), ("Timing status", 14),
                    ("Priced by", 11), ("Pricing date(s)", 16), ("Qty priced MT", 10),
                    ("Premium", 9), ("CBOT factor", 9), ("Priced CBOT", 10),
                    ("History on pricing day", 11), ("Window", 22), ("Best day (hindsight)", 12),
                    ("History on best day", 11), ("Highest in window", 11),
                    ("Best CBOT equivalent", 11), ("Opportunity, CBOT units", 11),
                    ("Opportunity USD/MT", 11), ("Opportunity USD", 13), ("FX", 9),
                    ("Opportunity EGP", 15), ("Opportunity % of CIF", 10),
                    ("Position in window (0 = low)", 11), ("20-day avg at the time", 11),
                    ("vs 20-day avg", 10), ("CBOT at arrival", 10), ("CBOT latest", 10),
                    ("Timing assessment", 15), ("Notes", 60)])
    for i, row in enumerate(rows):
        r = first + i
        t = row.get("timing") or {}
        ev = t.get("events") or []
        assessed = t.get("status") == "assessed"
        b.put(T, r, 1, row["id"]); b.put(T, r, 2, row["contract"]); b.put(T, r, 3, row["base"])
        b.put(T, r, 4, "assessed" if assessed else t.get("assessment"))
        b.put(T, r, 5, t.get("source") or "")
        dates = ", ".join(e["date"] for e in ev) if ev else (row.get("pricing_date") or "")
        b.put(T, r, 6, dates)
        b.put(T, r, 7, _blank(t.get("qty_priced")) if assessed else "", MONEY0)
        b.formula(T, r, 8, f'=IF({RG("premium", r)}="","",{RG("premium", r)})', row.get("premium"), MONEY2)
        b.formula(T, r, 9, f'=IF({RG("cbot_factor", r)}="","",{RG("cbot_factor", r)})',
                  row.get("cbot_factor"), "0.0000")
        ref = _weighted(ev, "history_on_day") if assessed else None
        lo = _weighted(ev, "best_history") if assessed else None
        hi = _weighted(ev, "worst_history") if assessed else None
        avg = _weighted(ev, "trailing_avg") if assessed else None
        b.put(T, r, 10, _blank(t.get("actual_cbot")) if assessed else "", MONEY2)
        b.put(T, r, 11, _blank(ref), MONEY2)
        b.put(T, r, 12, t.get("window") or "")
        b.put(T, r, 13, t.get("best_date") or "")
        b.put(T, r, 14, _blank(lo), MONEY2)
        b.put(T, r, 15, _blank(hi), MONEY2)
        b.formula(T, r, 16, f'=IF(J{r}="","",J{r}+(N{r}-K{r}))', t.get("best_cbot_equivalent"), MONEY2)
        opp_cbot = (ref - lo) if assessed else None
        b.formula(T, r, 17, f'=IF(J{r}="","",K{r}-N{r})', opp_cbot, MONEY2)
        b.formula(T, r, 18, f'=IF(Q{r}="","",Q{r}*I{r})', t.get("opportunity_usd_mt"), MONEY2)
        b.formula(T, r, 19, f'=IF(R{r}="","",R{r}*G{r})', t.get("opportunity_usd"), MONEY0)
        b.formula(T, r, 20, f'=IF({RG("fx", r)}="","",{RG("fx", r)})', row.get("fx"), "0.0000")
        b.formula(T, r, 21, f'=IF(OR(S{r}="",T{r}=""),"",S{r}*T{r})', t.get("opportunity_egp"), MONEY0)
        pct = t.get("opportunity_pct")
        b.formula(T, r, 22, f'=IF(OR(R{r}="",{RG("cif_usd_mt", r)}=""),"",R{r}/{RG("cif_usd_mt", r)})',
                  (pct / 100.0) if pct is not None else None, PCT)
        if assessed and len([e for e in ev if e.get("status") == "assessed"]) == 1:
            b.formula(T, r, 23, f'=IF(J{r}="","",IF(O{r}>N{r},(K{r}-N{r})/(O{r}-N{r}),0))',
                      t.get("range_position"), "0.00")
        else:
            b.put(T, r, 23, _blank(t.get("range_position")) if assessed else "", "0.00")
        b.put(T, r, 24, _blank(avg), MONEY2)
        b.formula(T, r, 25, f'=IF(J{r}="","",K{r}-X{r})', t.get("vs_trailing_avg"), "+#,##0.00;-#,##0.00")
        b.put(T, r, 26, _blank(row.get("cbot_at_delivery")), MONEY2)
        b.put(T, r, 27, _blank(row.get("cbot_latest")), MONEY2)
        b.formula(T, r, 28, f'=IF(D{r}<>"assessed",D{r},IF(V{r}*100<={small},"Favorable",'
                            f'IF(W{r}<=1/3,"Favorable",IF(W{r}<=2/3,"Acceptable","Unfavorable"))))',
                  t.get("assessment"))
        notes = t.get("reason") or ""
        if assessed:
            bits = [f"confidence: {t.get('confidence')}"]
            if t.get("spikes_excluded"):
                bits.append(f"{t['spikes_excluded']} one-day outlier(s) in the window not used as best/worst")
            if len(ev) > 1:
                bits.append("lots: figures weighted by tonnes")
            notes = "; ".join(bits)
        elif t.get("assessment") == tr.PENDING and t.get("forward"):
            fw = t["forward"]
            notes = (f"forward-looking only: CBOT {fw['current']:,.2f} on {fw['current_date']} sits at "
                     f"{fw['range_position']:.0%} of its last {fw['obs']} trading days' range "
                     f"({fw['low']:,.2f}–{fw['high']:,.2f}); not a forecast")
        b.put(T, r, 29, notes)
    b.put(T, tot, 1, "TOTAL", bold=True)
    b.formula(T, tot, 19, f"=SUM(S{first}:S{last})", model["summary"]["timing_opportunity_usd"],
              MONEY0, bold=True)
    b.formula(T, tot, 21, f"=SUM(U{first}:U{last})", model["summary"]["timing_opportunity_egp"],
              MONEY0, bold=True)
    b.verdict_colours(T, f"AB{first}:AB{last}")
    T.freeze_panes = "D5"
    T.auto_filter.ref = f"A4:AC{last}"

    # ── Data & Methodology (market data first: charts read it) ───────────
    _methodology(b, Dm, model)
    _market_and_points(b, Dm, T, model)

    # ── Executive Summary: decision table (final verdicts) ───────────────
    dec_top = 46
    dec_first = dec_top + 1
    b.put(E, dec_top - 1, 1, "CONTRACT DECISIONS — every contract, its own result", bold=True, size=12, color=NAVY)
    b.header(E, dec_top, [("ID", 8), ("Contract", 24), ("Base", 8), ("Supplier", 16), ("Qty MT", 10),
                          ("Saving vs local EGP", 15), ("Local result", 13),
                          ("Timing opportunity USD (hindsight)", 15), ("Timing assessment", 14),
                          ("Arrival verdict", 24), ("Local score", 7), ("Timing score", 7),
                          ("Final assessment", 15), ("Why", 70)])
    for i, row in enumerate(rows):
        r = dec_first + i
        rr = first + i
        b.put(E, r, 1, row["id"]); b.put(E, r, 2, row["contract"]); b.put(E, r, 3, row["base"])
        b.put(E, r, 4, row["supplier"])
        b.formula(E, r, 5, f"={RG('qty_mt', rr)}", row.get("qty_mt") or 0, MONEY0)
        b.formula(E, r, 6, f"={sav}!M{rr}", row.get("saving_local_egp"), MONEY0)
        b.formula(E, r, 7, f"={sav}!P{rr}", row.get("local_assessment"))
        b.formula(E, r, 8, f"={_q('CBOT & Timing')}!S{rr}", (row.get("timing") or {}).get("opportunity_usd"), MONEY0)
        b.formula(E, r, 9, f"={_q('CBOT & Timing')}!AB{rr}", (row.get("timing") or {}).get("assessment"))
        b.formula(E, r, 10, f"={con}!J{c_first + i}", (row.get("coverage") or {}).get("verdict"))
        score = lambda v: 1 if v == tr.FAVORABLE else (-1 if v == tr.UNFAVORABLE else 0)
        b.formula(E, r, 11, f'=IF(G{r}="Favorable",1,IF(G{r}="Unfavorable",-1,0))',
                  score(row.get("local_assessment")))
        b.formula(E, r, 12, f'=IF(I{r}="Favorable",1,IF(I{r}="Unfavorable",-1,0))',
                  score((row.get("timing") or {}).get("assessment")))
        b.formula(E, r, 13, f'=IF(G{r}="Insufficient Data","Insufficient Data",'
                            f'IF(OR(I{r}="Favorable",I{r}="Acceptable",I{r}="Unfavorable"),'
                            f'IF(K{r}+L{r}>=1,"Favorable",IF(K{r}+L{r}<=-1,"Unfavorable","Acceptable")),G{r}))',
                  row.get("final_assessment"))
        b.put(E, r, 14, _why(row))
    dec_last = dec_first + n - 1
    b.verdict_colours(E, f"G{dec_first}:G{dec_last}")
    b.verdict_colours(E, f"I{dec_first}:I{dec_last}")
    b.verdict_colours(E, f"M{dec_first}:M{dec_last}")
    E.auto_filter.ref = f"A{dec_top}:N{dec_last}"

    _executive(b, E, model, first, last, dec_first, dec_last, c_first)
    _missed(b, M, model, rows, first)
    _simulator_table(b, X, rows, first, RG, sav)

    b.wb.calculation.fullCalcOnLoad = True
    b.wb.save(path)
    return store_formula_results(path, b.cache)


def _why(row):
    t = row.get("timing") or {}
    cov = row.get("coverage") or {}
    bits = []
    if row.get("saving_local_egp_mt") is not None:
        bits.append(f"{row['saving_local_egp_mt']:+,.0f} EGP/MT vs local ({row.get('saving_kind')})")
    else:
        bits.append("no local comparison")
    if t.get("status") == "assessed":
        bits.append(f"CBOT fixed at {t['range_position']:.0%} of its window; hindsight best "
                    f"{t['best_date']} ({t['opportunity_usd_mt']:,.2f} USD/MT)")
    else:
        bits.append(f"timing: {t.get('reason') or t.get('assessment')}")
    if cov.get("cover_days") is not None:
        bits.append(f"{cov['cover_days']:,.1f} days of stock at arrival")
    return "; ".join(bits)


def _executive(b, E, model, first, last, dec_first, dec_last, c_first):
    s = model["summary"]
    reg, sav, tim = _q("Contract Register"), _q("Savings vs Local"), _q("CBOT & Timing")
    b.title(E, model.get("title") or "Procurement Decision & Contract Timing Report",
            f"Generated {model.get('generated')} from the app's data · as of {model.get('as_of')} · "
            f"{s['contracts']} contracts · savings vs local and timing are separate measures and are never added together")
    cards = [
        ("Total contract value, USD", f"=SUM({reg}!R{first}:R{last})", s["contract_value_usd"], MONEY0),
        ("Total contract value, EGP", f"=SUM({reg}!S{first}:S{last})", s["contract_value_egp"], MONEY0),
        ("Saving vs local — realized, EGP", f'=SUMIFS({sav}!M{first}:M{last},{sav}!O{first}:O{last},"Realized")',
         s["saving_local_realized_egp"], MONEY0),
        ("Saving vs local — estimated (open), EGP", f'=SUMIFS({sav}!M{first}:M{last},{sav}!O{first}:O{last},"Estimated")',
         s["saving_local_estimated_egp"], MONEY0),
        ("Timing opportunity (hindsight), USD", f"=SUM({tim}!S{first}:S{last})", s["timing_opportunity_usd"], MONEY0),
        ("Timing opportunity (hindsight), EGP", f"=SUM({tim}!U{first}:U{last})", s["timing_opportunity_egp"], MONEY0),
        ("Contracts analysed", f"=COUNTA({reg}!A{first}:A{last})", s["contracts"], INT),
        ("Timing assessed", f'=COUNTIF({tim}!D{first}:D{last},"assessed")', s["timing_assessed"], INT),
        ("Favorable timing", f'=COUNTIF({tim}!AB{first}:AB{last},"Favorable")', s["timing_favorable"], INT),
        ("Potentially missed savings (unfavorable timing)", f'=COUNTIF({tim}!AB{first}:AB{last},"Unfavorable")',
         s["missed_savings"], INT),
        ("Contracts needing review", f'=SUMPRODUCT(((M{dec_first}:M{dec_last}="Unfavorable")+(M{dec_first}:M{dec_last}='
                                     f'"Insufficient Data")+(I{dec_first}:I{dec_last}="Insufficient Data")>0)*1)',
         s["needs_review"], INT),
        ("Timing: insufficient data", f'=COUNTIF({tim}!AB{first}:AB{last},"Insufficient Data")',
         s["timing_insufficient"], INT),
    ]
    for k, (label, formula, value, fmt) in enumerate(cards):
        r = 5 + (k // 4) * 3
        col = 1 + (k % 4) * 3
        b.put(E, r, col, label, bold=True, color="FFFFFF", fill=BLUE, wrap=True)
        E.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
        b.formula(E, r + 1, col, formula, value, fmt, bold=True)
        E.cell(r + 1, col).font = Font(bold=True, size=14, color=NAVY)
        E.cell(r + 1, col).fill = PatternFill("solid", fgColor="EEF3F9")
        E.merge_cells(start_row=r + 1, start_column=col, end_row=r + 1, end_column=col + 1)
        E.row_dimensions[r].height = 30
    for col in range(1, 13):
        E.column_dimensions[get_column_letter(col)].width = max(
            E.column_dimensions[get_column_letter(col)].width or 0, 13)

    # Small summary tables the charts read
    r0 = 15
    b.put(E, r0, 1, "SAVING VS LOCAL BY COMMODITY, EGP", bold=True)
    b.put(E, r0 + 1, 1, "Base", bold=True); b.put(E, r0 + 1, 2, "Realized", bold=True)
    b.put(E, r0 + 1, 3, "Estimated", bold=True)
    bases = sorted({r["base"] for r in model["contracts"]})
    for i, base in enumerate(bases):
        r = r0 + 2 + i
        b.put(E, r, 1, base)
        for col, kind in ((2, "Realized"), (3, "Estimated")):
            v = sum(x["saving_local_egp"] for x in model["contracts"]
                    if x["base"] == base and x.get("saving_kind") == kind and x.get("saving_local_egp") is not None)
            b.formula(E, r, col, f'=SUMIFS({sav}!M{first}:M{last},{sav}!C{first}:C{last},A{r},'
                                 f'{sav}!O{first}:O{last},"{kind}")', v, MONEY0)
    chart = BarChart()
    chart.title, chart.height, chart.width = "Saving vs local by commodity, EGP", 7, 13
    chart.add_data(Reference(E, min_col=2, max_col=3, min_row=r0 + 1, max_row=r0 + 1 + len(bases)),
                   titles_from_data=True)
    chart.set_categories(Reference(E, min_col=1, min_row=r0 + 2, max_row=r0 + 1 + len(bases)))
    chart.y_axis.number_format = '#,##0.0,,"M"'
    chart.x_axis.delete = False; chart.y_axis.delete = False
    E.add_chart(chart, "E15")

    v0 = r0 + 3 + len(bases)
    b.put(E, v0, 1, "FINAL ASSESSMENT COUNT", bold=True)
    for i, word in enumerate((tr.FAVORABLE, tr.ACCEPTABLE, tr.UNFAVORABLE, tr.INSUFFICIENT)):
        r = v0 + 1 + i
        b.put(E, r, 1, word)
        b.formula(E, r, 2, f'=COUNTIF(M{dec_first}:M{dec_last},A{r})',
                  s["final_counts"].get(word, 0), INT)
    vc = BarChart()
    vc.title, vc.height, vc.width, vc.legend = "Final assessment", 7, 9, None
    vc.add_data(Reference(E, min_col=2, min_row=v0 + 1, max_row=v0 + 4))
    vc.set_categories(Reference(E, min_col=1, min_row=v0 + 1, max_row=v0 + 4))
    vc.x_axis.delete = False; vc.y_axis.delete = False
    E.add_chart(vc, "K15")

    f0 = 33
    b.put(E, f0, 1, "MAJOR FINDINGS", bold=True, size=12, color=NAVY)
    for i, line in enumerate(_findings(model)):
        b.put(E, f0 + 1 + i, 1, f"•  {line}")


def _findings(model):
    s, rows = model["summary"], model["contracts"]
    by_id = {r["id"]: r for r in rows}
    out = [f"Saving against local: {s['saving_local_realized_egp']:+,.0f} EGP realized on closed, "
           f"priced contracts, plus {s['saving_local_estimated_egp']:+,.0f} EGP estimated on open ones "
           f"(marked to today). Reported apart from timing."]
    if s["ranking"]:
        top = by_id[s["ranking"][0]]
        t = top["timing"]
        fixed = ", ".join(e["date"] for e in t.get("events") or [] if e.get("status") == "assessed")
        out.append(f"Largest hindsight timing opportunity: {top['contract']} ({top['id']}), CBOT fixed "
                   f"{fixed}; the lowest CBOT in its window was on {t['best_date']}: "
                   f"{t['opportunity_usd']:,.0f} USD ({t['opportunity_pct']:.1f}% of CIF). Hindsight only.")
    out.append(f"Timing assessed for {s['timing_assessed']} of {s['contracts']} contracts: "
               f"{s['timing_favorable']} favorable, {s['timing_acceptable']} acceptable, "
               f"{s['timing_unfavorable']} unfavorable. {s['timing_insufficient']} have a premium but no "
               f"pricing date (Insufficient Data); {s['timing_not_applicable']} are flat-price "
               f"(no CBOT to time); {s['timing_pending']} are open and still to be priced.")
    tight = [r for r in rows if str((r.get("coverage") or {}).get("verdict", "")).startswith("Needed")]
    early = [r for r in rows if str((r.get("coverage") or {}).get("verdict", "")).startswith("Early")]
    out.append(f"{len(tight)} arrivals came when stock was below the safety cover — buying them "
               f"earlier prevented (or shortened) a stoppage; {len(early)} arrived with ample stock.")
    for st in model.get("stock") or []:
        if st.get("shortage_mt"):
            out.append(f"{st['base']}: {st['shortage_mt']:,.0f} MT of consumption at "
                       f"{st['daily_rate']:,.0f} MT/day had no recorded cargo to come from — "
                       f"opening stock missing, or the rate is above what was bought.")
    if model.get("gaps"):
        kinds = {}
        for g in model["gaps"]:
            kinds[g["missing"]] = kinds.get(g["missing"], 0) + 1
        out.append("Missing data: " + ", ".join(f"{k} on {v} contract(s)" for k, v in sorted(kinds.items()))
                   + ". Signing date, shipment date and lead time are not recorded for any contract.")
    return out


def _missed(b, M, model, rows, first):
    tim, sav = _q("CBOT & Timing"), _q("Savings vs Local")
    b.title(M, "Missed Opportunities — ranked by hindsight timing opportunity",
            "Each figure is what fixing the CBOT on the window's lowest day would have saved. "
            "Hindsight benchmark, not a guaranteed opportunity. The saving vs local is shown "
            "beside it as a separate measure and is not added to it.")
    b.header(M, 4, [("Rank", 6), ("ID", 8), ("Contract", 24), ("Base", 8), ("Pricing date(s)", 16),
                    ("Best day (hindsight)", 12), ("Opportunity USD", 14), ("Opportunity EGP", 15),
                    ("Opportunity % of CIF", 10), ("Timing assessment", 14),
                    ("Saving vs local EGP (separate)", 15), ("Confidence", 40)])
    index = {r["id"]: i for i, r in enumerate(rows)}
    ranking = model["summary"]["ranking"]
    for k, cid in enumerate(ranking):
        r, rr = 5 + k, first + index[cid]
        row = rows[index[cid]]
        t = row["timing"]
        b.put(M, r, 1, k + 1); b.put(M, r, 2, cid); b.put(M, r, 3, row["contract"])
        b.put(M, r, 4, row["base"])
        b.formula(M, r, 5, f"={tim}!F{rr}", ", ".join(e["date"] for e in t.get("events") or []))
        b.formula(M, r, 6, f"={tim}!M{rr}", t.get("best_date"))
        b.formula(M, r, 7, f"={tim}!S{rr}", t.get("opportunity_usd"), MONEY0)
        b.formula(M, r, 8, f"={tim}!U{rr}", t.get("opportunity_egp"), MONEY0)
        b.formula(M, r, 9, f"={tim}!V{rr}", (t.get("opportunity_pct") or 0) / 100.0, PCT)
        b.formula(M, r, 10, f"={tim}!AB{rr}", t.get("assessment"))
        b.formula(M, r, 11, f"={sav}!M{rr}", row.get("saving_local_egp"), MONEY0)
        b.put(M, r, 12, t.get("confidence") or "")
    end = 4 + len(ranking)
    b.verdict_colours(M, f"J5:J{max(end, 5)}")
    if ranking:
        top_n = min(15, len(ranking))
        ch = BarChart()
        ch.type, ch.title, ch.height, ch.width, ch.legend = "bar", "Top timing opportunities, USD", 9, 16, None
        ch.add_data(Reference(M, min_col=7, min_row=5, max_row=4 + top_n))
        ch.set_categories(Reference(M, min_col=3, min_row=5, max_row=4 + top_n))
        ch.y_axis.number_format = "#,##0"
        ch.x_axis.scaling.orientation = "maxMin"
        ch.x_axis.delete = False; ch.y_axis.delete = False
        M.add_chart(ch, "N4")
    # Arrivals where buying earlier kept the plant running
    s0 = end + 3
    b.put(M, s0, 1, "ARRIVALS THAT PREVENTED A SHORTAGE — earlier buying justified by stock", bold=True)
    # Same columns as the ranking above (B ID, C Contract, D Base…), so the
    # widths that suit one suit the other.
    b.header(M, s0 + 1, [("ID", 8), ("Contract", 24), ("Base", 8), ("Arrival", 16),
                         ("Cover before arrival, days", 12), ("Verdict", 14)], start=2, height=30)
    tight = [r for r in rows if str((r.get("coverage") or {}).get("verdict", "")).startswith("Needed")]
    for k, row in enumerate(tight):
        r = s0 + 2 + k
        b.put(M, r, 2, row["id"]); b.put(M, r, 3, row["contract"]); b.put(M, r, 4, row["base"])
        b.put(M, r, 5, row.get("delivery_date"))
        b.put(M, r, 6, round(row["coverage"]["cover_days"], 1), "0.0")
        b.put(M, r, 7, row["coverage"]["verdict"])
    M.column_dimensions["G"].width = 28
    M.freeze_panes = "D5"


def _simulator_table(b, X, rows, first, RG, sav):
    top = 16
    b.put(X, top - 1, 1, "SCENARIO RESULTS — every contract under the assumptions above", bold=True)
    b.header(X, top, [("ID", 8), ("Contract", 58), ("Base", 8), ("Qty MT", 10), ("Base CIF USD/MT", 11),
                      ("CBOT factor", 9), ("Scenario CIF USD/MT", 11), ("Base FX", 9),
                      ("Scenario FX", 9), ("Fees EGP/MT", 10), ("Scenario landed EGP/MT", 12),
                      ("Base local EGP/MT", 11), ("Scenario local EGP/MT", 11),
                      ("Scenario saving EGP/MT", 12), ("Scenario saving EGP", 14),
                      ("Base saving EGP", 14), ("Change EGP", 13)])
    s_first = top + 1
    for i, row in enumerate(rows):
        r, rr = s_first + i, first + i
        b.put(X, r, 1, row["id"]); b.put(X, r, 2, row["contract"]); b.put(X, r, 3, row["base"])
        b.formula(X, r, 4, f'=IF({RG("qty_mt", rr)}="","",{RG("qty_mt", rr)})', row.get("qty_mt"), MONEY0)
        b.formula(X, r, 5, f'=IF({RG("cif_usd_mt", rr)}="","",{RG("cif_usd_mt", rr)})', row.get("cif_usd_mt"), MONEY2)
        b.formula(X, r, 6, f'=IF({RG("cbot_factor", rr)}="","",{RG("cbot_factor", rr)})',
                  row.get("cbot_factor"), "0.0000")
        # Only a premium contract moves with the CBOT; a flat price is fixed.
        b.formula(X, r, 7, f'=IF(E{r}="","",IF(OR(F{r}="",N({RG("premium", rr)})=0),E{r},'
                           f'E{r}+($C$4+$C$5)*F{r}))', row.get("cif_usd_mt"), MONEY2)
        b.formula(X, r, 8, f'=IF({RG("fx", rr)}="","",{RG("fx", rr)})', row.get("fx"), "0.0000")
        b.formula(X, r, 9, f'=IF(H{r}="","",H{r}*(1+$C$6/100))', row.get("fx"), "0.0000")
        fees = sum(row.get(k) or 0.0 for k in ("discharge_egp_mt", "clearance_egp_mt", "freight_egp_mt"))
        b.formula(X, r, 10, f'=N({RG("discharge_egp_mt", rr)})+N({RG("clearance_egp_mt", rr)})'
                            f'+N({RG("freight_egp_mt", rr)})', fees, MONEY2)
        b.formula(X, r, 11, f'=IF(OR(G{r}="",I{r}=""),"",G{r}*I{r}+J{r})', row.get("landed_cost_egp_mt"), MONEY2)
        b.formula(X, r, 12, f'=IF({RG("local_egp_mt", rr)}="","",{RG("local_egp_mt", rr)})',
                  row.get("local_egp_mt"), MONEY2)
        b.formula(X, r, 13, f'=IF(L{r}="","",L{r}*(1+$C$7/100))', row.get("local_egp_mt"), MONEY2)
        b.formula(X, r, 14, f'=IF(OR(K{r}="",M{r}=""),"",M{r}-K{r})', row.get("saving_local_egp_mt"), MONEY2)
        b.formula(X, r, 15, f'=IF(OR(N{r}="",D{r}=""),"",N{r}*D{r})', row.get("saving_local_egp"), MONEY0)
        b.formula(X, r, 16, f"={sav}!M{rr}", row.get("saving_local_egp"), MONEY0)
        base = row.get("saving_local_egp")
        b.formula(X, r, 17, f'=IF(OR(O{r}="",P{r}=""),"",O{r}-P{r})', 0.0 if base is not None else None, MONEY0)
    s_last = s_first + len(rows) - 1
    total = sum(r["saving_local_egp"] for r in rows if r.get("saving_local_egp") is not None)
    b.put(X, s_last + 1, 1, "TOTAL", bold=True)
    b.formula(X, s_last + 1, 15, f"=SUM(O{s_first}:O{s_last})", total, MONEY0, bold=True)
    b.formula(X, s_last + 1, 16, f"=SUM(P{s_first}:P{s_last})", total, MONEY0, bold=True)
    b.formula(X, s_last + 1, 17, f"=SUM(Q{s_first}:Q{s_last})", 0.0, MONEY0, bold=True)
    X.freeze_panes = X.cell(s_first, 3)


def _methodology(b, D, model):
    p = model["params"]
    b.title(D, "Data & Methodology", "How every figure in this report is built, where the "
            "data comes from, and what is missing.")
    lines = [
        ("CALCULATIONS", None),
        ("Contract value", "Qty × CIF (USD); × FX for EGP."),
        ("FX", "The contract's cost FX: Form 4 FX once entered (every commodity except SBM), else the "
               "delivery FX, else today's market FX — the app's own rule."),
        ("Landed cost", "CIF × FX + intake/discharge + clearance + freight (EGP/MT)."),
        ("Local price", "The local price logged on or before the arrival day, matched to the cargo's "
                        "origin (CORN-BRZ, CORN-ARG… else CORN), plus its transport — delivered, "
                        "like the landed cost."),
        ("Saving vs local", "Local − landed cost, per MT and × qty. Realized = closed and marked priced "
                            "(the Savings Tracker's figure); Estimated = open, marked to today."),
        ("Priced CBOT", "From the contract's own numbers: CIF ÷ CBOT factor − premium; or the CBOT stored "
                        "on each pricing lot. Factors: corn 0.3937, soybean/wheat 0.36745, SBM 1.1023."),
        ("Timing window", f"{p['window_before_days']} days before to {p['window_after_days']} days after each "
                          "pricing day, never past the data's last day."),
        ("Best CBOT equivalent", "Priced CBOT + (stored history on the best day − stored history on the "
                                 "pricing day). Only the history's moves are used, because it is the "
                                 "front-month series while the contract is priced on its own futures month."),
        ("Timing opportunity", "(Priced CBOT − best CBOT equivalent) × factor × tonnes priced, in USD; "
                               "× FX in EGP; ÷ CIF for %. Equal to (actual comparable cost − best feasible "
                               "alternative cost) × quantity, with premium, fees and FX held as contracted."),
        ("Position in window", "(history on pricing day − lowest) ÷ (highest − lowest): 0 = fixed at the low."),
        ("Timing assessment", f"Favorable if the opportunity ≤ {p['small_opportunity_pct']}% of CIF or the "
                              "position ≤ 1/3; Acceptable ≤ 2/3; else Unfavorable."),
        ("Local result", f"Favorable if the saving ≥ +{p['local_band_pct']}% of local; Unfavorable if "
                         f"≤ −{p['local_band_pct']}%; else Acceptable."),
        ("Final assessment", "Scores Favorable +1, Acceptable 0, Unfavorable −1, local + timing: ≥ +1 "
                             "Favorable, 0 Acceptable, ≤ −1 Unfavorable. Without a timing result the local "
                             "result stands; without a local result: Insufficient Data."),
        ("Stock at arrival", "The app's FIFO walk at the configured daily rate from each delivery; stock "
                             "the day before arrival ÷ daily rate = cover days."),
        ("HINDSIGHT AND FORESIGHT", None),
        ("Hindsight", "The best day uses prices after the decision. It is a benchmark for learning, not a "
                      "claim the day was knowable or achievable."),
        ("At the time", f"The 20-day columns use only the {p['trailing_obs']} trading days on or before the "
                        "pricing day — what a buyer could have seen."),
        ("Open contracts", "Still to be priced: only where today's CBOT sits in its recent range — "
                           "forward-looking context, not a forecast."),
        ("LIMITATIONS", None),
        ("Spread", "Using front-month moves assumes the spread to the priced month held across the window; "
                   "a futures roll inside the window (typically 5–20 ¢ for corn) can shift a result."),
        ("Outliers", f"Days that jump ≥ {p['spike_pct']}% away from both neighbours and come back are "
                     "treated as stray prints: never offered as best or worst; a pricing day that is one "
                     "lowers the result's confidence."),
        ("Feasibility", "Moving the pricing day does not move the shipment, so it carries no stock risk of "
                        "its own; the contract is assumed to have allowed pricing anywhere in the window. "
                        "Signing date, shipment date and lead time are not in the data, so the timing of "
                        "the purchase itself (not its pricing) cannot be evaluated."),
        ("Nothing invented", "No CBOT, FX, premium or local price is estimated. A missing input is listed "
                             "below and the result is marked Insufficient Data."),
    ]
    r = 4
    for label, text in lines:
        if text is None:
            r += 1
            b.put(D, r, 1, label, bold=True, color=NAVY)
        else:
            b.put(D, r, 1, label, bold=True)
            b.put(D, r, 2, text, wrap=True)
        r += 1
    D.column_dimensions["A"].width = 24
    D.column_dimensions["B"].width = 100
    r += 1
    b.put(D, r, 1, "DATA SOURCES", bold=True, color=NAVY)
    r += 1
    for name, desc in (("contracts", f"{len(model['contracts'])} contract records"),
                       ("cbot_history", "; ".join(f"{k}: {len(v)} days {v[0][0]} … {v[-1][0]}"
                                                  for k, v in model["market"].items() if v)),
                       ("fx_history", f"{len(model.get('fx') or [])} days"),
                       ("local_prices", "delivered local prices by commodity key"),
                       ("consumption", "daily consumption rate by commodity")):
        b.put(D, r, 1, name, bold=True); b.put(D, r, 2, desc); r += 1
    r += 1
    b.put(D, r, 1, "MISSING INFORMATION", bold=True, color=NAVY)
    r += 1
    b.put(D, r, 1, "ID", bold=True); b.put(D, r, 2, "Contract — missing", bold=True)
    for g in model.get("gaps") or []:
        r += 1
        b.put(D, r, 1, g["id"]); b.put(D, r, 2, f"{g['contract']} — {g['missing']}")
    r += 2
    b.put(D, r, 1, "PRICING LOTS", bold=True, color=NAVY)
    for row in model["contracts"]:
        for e in (row.get("timing") or {}).get("events") or []:
            if (row.get("timing") or {}).get("source") == "pricing lots":
                r += 1
                b.put(D, r, 1, row["id"])
                b.put(D, r, 2, f"{row['contract']}: {e['date']} · {e['qty_mt']:,.2f} MT at "
                               f"{e['priced_cbot']:,.2f} · history {e.get('history_on_day', '')}")


def _market_and_points(b, D, T, model):
    """Market history and the pricing/best points, plus a chart per commodity."""
    col = 4
    charts = 0
    for base, series in model["market"].items():
        if not series:
            continue
        b.put(D, 3, col, f"CBOT {base}", bold=True)
        b.put(D, 4, col, "Date", bold=True); b.put(D, 4, col + 1, "Price", bold=True)
        for i, (d, v) in enumerate(series):
            b.put(D, 5 + i, col, dt.date.fromisoformat(d), "yyyy-mm-dd")
            b.put(D, 5 + i, col + 1, v, MONEY2)
        n = len(series)
        pts = [(r["timing"]["events"], r) for r in model["contracts"]
               if r["base"] == base and (r.get("timing") or {}).get("status") == "assessed"]
        pc = col + 3
        b.put(D, 3, pc, f"{base} pricing and best days", bold=True)
        for j, h in enumerate(("Priced on", "History", "Best day", "History")):
            b.put(D, 4, pc + j, h, bold=True)
        k = 0
        for events, row in pts:
            for e in events:
                if e.get("status") != "assessed":
                    continue
                b.put(D, 5 + k, pc, dt.date.fromisoformat(e["history_date"]), "yyyy-mm-dd")
                b.put(D, 5 + k, pc + 1, e["history_on_day"], MONEY2)
                b.put(D, 5 + k, pc + 2, dt.date.fromisoformat(e["best_date"]), "yyyy-mm-dd")
                b.put(D, 5 + k, pc + 3, e["best_history"], MONEY2)
                k += 1
        for cc in range(col, pc + 4):
            D.column_dimensions[get_column_letter(cc)].width = 12
        sc = ScatterChart()
        sc.title = f"CBOT {base}: when contracts were priced vs the hindsight best day"
        sc.style, sc.height, sc.width = 13, 8, 22
        sc.x_axis.number_format = "mmm-yy"
        sc.x_axis.title, sc.y_axis.title = "Date", "CBOT"
        sc.x_axis.delete = False; sc.y_axis.delete = False
        xs = Reference(D, min_col=col, min_row=5, max_row=4 + n)
        line = ChartSeries(Reference(D, min_col=col + 1, min_row=5, max_row=4 + n), xs, title="CBOT history")
        line.marker.symbol = "none"
        line.smooth = False
        line.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="1F3A5F", w=15000))
        sc.series.append(line)
        if k:
            for (xc, yc, label, colour, symbol) in ((pc, pc + 1, "Priced", "C0392B", "circle"),
                                                    (pc + 2, pc + 3, "Best day (hindsight)", "1A7A1A", "triangle")):
                s = ChartSeries(Reference(D, min_col=yc, min_row=5, max_row=4 + k),
                                Reference(D, min_col=xc, min_row=5, max_row=4 + k), title=label)
                s.marker.symbol = symbol
                s.marker.size = 8
                s.marker.graphicalProperties = GraphicalProperties(
                    solidFill=colour, ln=LineProperties(solidFill=colour))
                s.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
                sc.series.append(s)
        T.add_chart(sc, f"AE{4 + charts * 18}")
        charts += 1
        col = pc + 6
