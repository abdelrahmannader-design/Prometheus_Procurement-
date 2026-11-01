"""The management savings workbook — structured tables, live formulas.

Built from the model in ``mgmt_report``. Contract data, local market
observations, consumption allocations and each comparison sit in their
own Excel tables; every derived figure is a formula over the rows it comes
from (daily rows → contract summaries → management summary), and the
value of every formula is stored beside it so viewers that do not
calculate still show the numbers.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .mgmt_report import MATCHED
from .xlsx_cache import store_formula_results

SHEETS = ("Management Summary", "Contract Master", "Local Market Price History",
          "Daily Consumption Allocation", "Savings vs Local", "Contract Savings Summary",
          "CBOT Fixed vs Consumption", "Contract CBOT Performance",
          "Open Contracts vs Local Today", "Excluded & Missing Data", "Validation & Sources",
          "Local Purchases", "Monthly Trend")
NAVY, GREEN, RED, GREY, AMBER = "1F3A5F", "1A7A1A", "B3261E", "6B7280", "9A6B00"
M0, M2, PCT = "#,##0;[Red]-#,##0", "#,##0.00;[Red]-#,##0.00", "0.0%"
NA = "\"N/A — Missing Source Data\""  # Excel string literal for unsupported calculations
FIRST = 5  # first data row of every table (header on row 4)


def q(title):
    return "'" + title.replace("'", "''") + "'"


def blank(v):
    return "" if v is None else v


class Book:
    def __init__(self):
        self.wb = Workbook()
        self.wb.active.title = SHEETS[0]
        for s in SHEETS[1:]:
            self.wb.create_sheet(s)
        self.cache = {s: {} for s in SHEETS}

    def ws(self, i):
        return self.wb[SHEETS[i]]

    def put(self, ws, r, c, v, fmt=None, bold=False, color=None, fill=None, size=None, wrap=False, italic=False):
        cell = ws.cell(r, c, v)
        if fmt:
            cell.number_format = fmt
        if bold or color or size or italic:
            cell.font = Font(bold=bold, color=color, size=size, italic=italic)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if wrap:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        return cell

    def f(self, ws, r, c, formula, value, fmt=None, **kw):
        self.put(ws, r, c, formula, fmt, **kw)
        self.cache[ws.title][f"{get_column_letter(c)}{r}"] = "" if value is None else value

    def title(self, ws, text, sub):
        self.put(ws, 1, 1, text, bold=True, size=15, color=NAVY)
        self.put(ws, 2, 1, sub, italic=True, color=GREY)

    def table(self, ws, name, heads, n_rows, widths=None):
        for i, h in enumerate(heads, start=1):
            ws.cell(4, i, h)
            if widths:
                ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
        last = FIRST + max(n_rows, 1) - 1
        ref = f"A4:{get_column_letter(len(heads))}{last}"
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
        ws.freeze_panes = "C5"
        ws.row_dimensions[4].height = 32
        for i in range(1, len(heads) + 1):
            ws.cell(4, i).alignment = Alignment(wrap_text=True, vertical="center")
        return last

    def signs(self, ws, ref):
        ws.conditional_formatting.add(ref, CellIsRule(operator="lessThan", formula=["0"], font=Font(color=RED)))
        ws.conditional_formatting.add(ref, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=GREEN)))


def place(b, name, after):
    """Put sheet ``name`` straight after sheet ``after`` in the workbook."""
    names = b.wb.sheetnames
    b.wb.move_sheet(name, offset=(names.index(after) + 1) - names.index(name)
                    if names.index(name) > names.index(after) else 0)


def col(sheet, letter, last):
    return f"{q(SHEETS[sheet])}!${letter}${FIRST}:${letter}${last}"


def build_workbook(model, path):
    b = Book()
    last = {}
    last["cm"] = _contract_master(b, model)
    last["lp"] = _local_history(b, model)
    last["al"] = _allocations(b, model)
    last["lq"] = _local_purchases(b, model)
    _contract_master_consumed(b, model, last)
    last["sv"] = _savings(b, model)
    last["cs"] = _contract_savings(b, model, last)
    last["cb"] = _cbot(b, model)
    last["cp"] = _cbot_performance(b, model, last)
    _monthly(b, model, last)
    last["op"] = _open(b, model)
    last["ex"] = _excluded(b, model)
    _validation(b, model, last)
    _summary(b, model, last)
    for ws in b.wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        if ws.title != SHEETS[0]:
            ws.print_title_rows = "4:4"
        ws.oddFooter.center.text = "&A · page &P of &N"
    b.wb.calculation.fullCalcOnLoad = True
    b.wb.save(path)
    return store_formula_results(path, b.cache)


# ── 1. Contract Master ───────────────────────────────────────────────────

def _contract_master(b, m):
    ws = b.ws(1)
    rows = m["contracts"]
    b.title(ws, "Contract Master", "One row per contract record in the database (exact duplicates removed). "
            "CBOT Fixed: recorded pricing lots, else implied = CIF ÷ factor − premium — shown even when the "
            "pricing date is not recorded. Unpriced contracts show Not fixed and a separate indicative value.")
    heads = ["Contract ID", "Contract Number", "Supplier", "Commodity", "Origin", "Contract Status",
             "Contract Quantity MT", "Delivered Quantity MT", "Outstanding Quantity MT", "Consumed Quantity MT",
             "Available Stock MT", "Delivery Date", "Delivery Status", "CIF USD/MT", "Premium", "Premium Unit",
             "Futures Month", "Conversion Factor", "CBOT Fixed", "CBOT Fixing Type", "Indicative CBOT (not fixed)",
             "Pricing Date", "Exchange Rate", "FX Basis", "Landed Cost Incl. Transport EGP/MT", "Pricing Status",
             "Source Record ID", "Fixing Note", "Base", "Storage Start Date", "Receipt Status",
             "Pricing Status Class", "Final CIF USD/MT", "Form 4 FX", "Form 4 Date", "Inland Transport EGP/MT",
             "Discharge EGP/MT", "Clearance EGP/MT", "Landed Cost Excl. Transport EGP/MT", "Reconciliation Flags"]
    widths = [9, 22, 16, 11, 11, 9, 10, 10, 11, 10, 10, 11, 24, 9, 9, 10, 14, 9, 10, 26, 11, 13, 9, 11, 11,
              10, 16, 40, 8, 12, 24, 13, 10, 9, 11, 11, 10, 10, 12, 50]
    last = b.table(ws, "ContractMaster", heads, len(rows), widths)
    for i, c in enumerate(rows):
        r = FIRST + i
        f, cm = c["fix"], c.get("contract_month") or {}
        vals = [c["id"], c["contract"], c.get("supplier"), c.get("commodity"), c.get("origin"), c.get("status"),
                c.get("qty_mt"), c["delivered_mt"]]
        for j, v in enumerate(vals, start=1):
            b.put(ws, r, j, blank(v), M0 if j >= 7 else None)
        b.f(ws, r, 9, f"=G{r}-H{r}", c["outstanding_mt"], M0)
        b.put(ws, r, 11, c["available_mt"], M0)
        b.put(ws, r, 12, c.get("delivery_date") or "")
        b.put(ws, r, 13, c["delivery_status"])
        b.put(ws, r, 14, blank(c.get("cif_usd_mt")), M2)
        b.put(ws, r, 15, blank(c.get("premium")), M2)
        b.put(ws, r, 16, f["unit"] if f["eligible"] else "")
        b.put(ws, r, 17, (cm.get("code", "") + (" (year inferred)" if cm.get("inferred") else ""))
              if cm else (c.get("futures_month") or ""))
        b.put(ws, r, 18, blank(f["factor"]) if f["eligible"] else "", "0.0000")
        if f["type"] == "Implied":
            b.f(ws, r, 19, f"=N{r}/R{r}-O{r}", f["value"], M2, bold=True)
        elif f["type"] == "Recorded":
            b.put(ws, r, 19, f["value"], M2, bold=True)
        elif f["type"] == "Not fixed":
            b.put(ws, r, 19, "Not fixed", color=AMBER)
        elif f["eligible"]:
            b.put(ws, r, 19, "Not assessable", color=GREY)
        b.put(ws, r, 20, f["status"])
        if f["type"] == "Not fixed" and f["indicative"] is not None:
            b.f(ws, r, 21, f"=N{r}/R{r}-O{r}", f["indicative"], M2, color=AMBER)
        b.put(ws, r, 22, f["pricing_date"] or ("Not recorded" if f["fixed"] else ""))
        b.put(ws, r, 23, blank(c.get("fx")), "0.0000")
        b.put(ws, r, 24, c.get("fx_basis") or "")
        b.put(ws, r, 25, blank(c.get("landed_egp_mt")), M2)
        b.put(ws, r, 26, c.get("pricing_status") or "")
        b.put(ws, r, 27, c.get("source_id") or "")
        b.put(ws, r, 28, f["note"])
        b.put(ws, r, 29, c["base"])
        nr = c["storage_status"] == "Not recorded"
        b.put(ws, r, 30, c["storage_status"], color=AMBER if nr else None, bold=nr)
        b.put(ws, r, 31, c["receipt_status"], color=AMBER if not c["received"] else None)
        b.put(ws, r, 32, c["pricing_class"])
        b.put(ws, r, 33, blank(c.get("cif_usd_mt")) if c["pricing_class"] == "Fixed" else "", M2)
        b.put(ws, r, 34, blank(c.get("form4_fx")), "0.0000")
        b.put(ws, r, 35, c.get("form4_date") or ("Not recorded" if not c.get("form4_fx") else ""))
        b.put(ws, r, 36, blank(c.get("freight_egp_mt")) if c.get("freight_egp_mt") is not None else "N/A — Missing Source Data", M2)
        b.put(ws, r, 37, blank(c.get("discharge_egp_mt")), M2)
        b.put(ws, r, 38, blank(c.get("clearance_egp_mt")), M2)
        b.f(ws, r, 39, f'=IF(AND(ISNUMBER(Y{r}),ISNUMBER(AJ{r})),Y{r}-AJ{r},{NA})',
            (c["landed_egp_mt"] - c["freight_egp_mt"]) if None not in (c.get("landed_egp_mt"), c.get("freight_egp_mt"))
            else "N/A — Missing Source Data", M2)
        b.put(ws, r, 40, "; ".join(c.get("flags") or []), color=AMBER if c.get("flags") else None)
    return last


def _contract_master_consumed(b, m, last):
    ws = b.ws(1)
    for i, c in enumerate(m["contracts"]):
        r = FIRST + i
        b.f(ws, r, 10, f"=SUMIFS({col(3, 'G', last['al'])},{col(3, 'B', last['al'])},A{r})", c["consumed_mt"], M0)


# ── 2. Local Market Price History ────────────────────────────────────────

def _local_history(b, m):
    ws = b.ws(2)
    rows = sorted(m["local_obs"], key=lambda o: (o["date"], o["key"]))
    b.title(ws, "Local Market Price History", "Only the local price observations stored in the database — "
            "no rows are created. Where a quote is carried forward to a later consumption day, that is shown "
            "on Savings vs Local with the quote's own date and age.")
    heads = ["Price Observation Date", "Commodity", "Origin / Grade", "Local Market Price EGP/MT",
             "Local Freight EGP/MT", "Total Local Delivered Price EGP/MT", "Supplier / Source",
             "Source Record ID", "Data Status"]
    last = b.table(ws, "LocalPriceHistory", heads, len(rows), [12, 10, 13, 12, 11, 13, 22, 17, 24])
    for i, o in enumerate(rows):
        r = FIRST + i
        g = o.get
        b.put(ws, r, 1, g("date")); b.put(ws, r, 2, g("base") or str(g("key") or "").split("-")[0])
        b.put(ws, r, 3, g("key"))
        b.put(ws, r, 4, blank(g("price")), M2); b.put(ws, r, 5, blank(g("transport")), M2)
        b.f(ws, r, 6, f'=IF(OR(D{r}="",E{r}=""),"",D{r}+E{r})', g("delivered"), M2)
        b.put(ws, r, 7, g("source") or ""); b.put(ws, r, 8, g("source_id") or "")
        b.put(ws, r, 9, (g("status") or "Observed") + (" (freight: default rate)" if g("transport_default_used") else ""))
    return last


# ── 3. Daily Consumption Allocation ──────────────────────────────────────

def _allocations(b, m):
    ws = b.ws(3)
    rows = [a for a in m["allocations"] if a["type"] == "Import"]
    b.title(ws, "Daily Consumption Allocation", "Every consumed imported tonne attributed to a real contract. "
            "Local purchases are on their own sheet; unsourced consumption is on Excluded & Missing Data.")
    heads = ["Consumption Date", "Contract ID", "Contract Number", "Supplier", "Commodity", "Origin",
             "Quantity Consumed MT", "Allocation Method", "Contract Remaining Balance MT", "Contract Source ID",
             "Consumption Source / Assumption", "Record Type", "Base"]
    last = b.table(ws, "Allocations", heads, len(rows), [12, 12, 22, 16, 11, 11, 10, 13, 12, 18, 40, 13, 8])
    for i, a in enumerate(rows):
        r = FIRST + i
        for j, v in enumerate([a["date"], a["contract_id"], a["contract"], a["supplier"], a["commodity"],
                               a["origin"], a["mt"], a["method"], a.get("remaining_mt"), a["source_id"],
                               a["assumption"], a["type"], a["base"]], start=1):
            b.put(ws, r, j, blank(v), M0 if j in (7, 9) else None)
    return last


def _local_purchases(b, m):
    ws = b.ws(11)
    place(b, ws.title, SHEETS[3])
    rows = [a for a in m["allocations"] if a["type"] != "Import"]
    b.title(ws, "Local Purchases", "Tonnes supplied by local purchases, kept apart: never part of the import "
            "savings. Each is set against the local market price + freight observed for its day "
            "(positive = bought below that day's market).")
    heads = ["Consumption Date", "Purchase ID", "Purchase", "Supplier", "Commodity", "Quantity Consumed MT",
             "Purchase Cost EGP/MT", "Local Market + Freight EGP/MT", "Market Price Date", "vs Market EGP/MT",
             "vs Market EGP", "Base", "Allocation Method", "Source Record ID"]
    last = b.table(ws, "LocalPurchases", heads, len(rows), [12, 16, 30, 22, 10, 10, 12, 13, 12, 11, 13, 8, 13, 22])
    for i, a in enumerate(rows):
        r = FIRST + i
        loc = a.get("local") or {}
        price = loc.get("price")
        for j, v in enumerate([a["date"], a["contract_id"], a["contract"], a["supplier"], a["commodity"], a["mt"],
                               a.get("cost"), price, loc.get("date") or ""], start=1):
            b.put(ws, r, j, blank(v), M0 if j == 6 else (M2 if j in (7, 8) else None))
        per = (price - a["cost"]) if None not in (price, a.get("cost")) else None
        b.f(ws, r, 10, f'=IF(OR(G{r}="",H{r}=""),"",H{r}-G{r})', per, M2)
        b.f(ws, r, 11, f'=IF(J{r}="","",J{r}*F{r})', per * a["mt"] if per is not None else None, M0)
        b.put(ws, r, 12, a["base"]); b.put(ws, r, 13, a["method"]); b.put(ws, r, 14, a["source_id"])
    b.signs(ws, f"J{FIRST}:K{last}")
    return last


# ── 4. Savings vs Local ──────────────────────────────────────────────────

def _savings(b, m):
    ws = b.ws(4)
    rows = m["savings"]
    b.title(ws, "Savings vs Local", "Daily Savings = (Local Delivered Price − Imported Landed Cost) × "
            "Consumed MT. Only rows with a real contract, a landed cost and a local price in force that day (the last "
            "price entered stays in force until it changes); the rest are on Excluded & Missing Data, not counted as zero.")
    heads = ["Contract Number", "Contract ID", "Supplier", "Commodity", "Consumption Date", "Consumed MT",
             "Imported Cost Incl. Transport EGP/MT", "Local Price Incl. Transport EGP/MT", "Local Price Observation Date",
             "Quote Age (days)", "Local Price Record ID", "Savings Incl. Transport EGP/MT", "Total Savings Incl. Transport EGP", "CIF Basis",
             "Allocation Method", "Data Quality Status", "Local Price Of", "Local x MT", "Base", "Month",
             "Import Inland Transport EGP/MT", "Imported Cost Excl. Transport EGP/MT", "Local Transport EGP/MT",
             "Local Price Excl. Transport EGP/MT", "Savings Excl. Transport EGP/MT", "Total Savings Excl. Transport EGP",
             "Local Price Age Status", "Storage Start Date", "Local Excl x MT", "MT With Excl",
             "CBOT That Day", "CBOT Price Date", "FX That Day EGP/USD", "FX Date", "Port Costs (Discharge + Clearance) EGP/MT",
             "CBOT Factor", "Local Equivalent Premium (CIF basis)", "Local Equivalent Premium (port costs not removed)",
             "Contract Premium", "Premium Gap (Local Equivalent − Contract)", "Premium Unit"]
    last = b.table(ws, "SavingsVsLocal", heads, len(rows),
                   [22, 9, 16, 11, 12, 9, 12, 12, 12, 8, 16, 11, 13, 13, 13, 40, 11, 14, 8, 9,
                    11, 12, 10, 12, 11, 13, 34, 12, 13, 9, 10, 11, 10, 11, 13, 8, 14, 14, 10, 14, 10])
    for i, s in enumerate(rows):
        r = FIRST + i
        for j, v in enumerate([s["contract"], s["contract_id"], s["supplier"], s["commodity"], s["date"],
                               s["mt"], s["cost"], s["local"], s["local_date"], s["local_age"],
                               s["local_source"]], start=1):
            b.put(ws, r, j, blank(v), M0 if j == 6 else (M2 if j in (7, 8) else None))
        b.f(ws, r, 12, f"=H{r}-G{r}", s["saving_mt"], M2)
        b.f(ws, r, 13, f"=L{r}*F{r}", s["saving"], M0)
        b.put(ws, r, 14, s["cif_basis"]); b.put(ws, r, 15, s["method"]); b.put(ws, r, 16, s["quality"])
        b.put(ws, r, 17, s["local_key"])
        b.f(ws, r, 18, f"=H{r}*F{r}", s["local"] * s["mt"], M0)
        b.put(ws, r, 19, s["base"])
        b.put(ws, r, 20, s["date"][:7])
        b.put(ws, r, 21, s["freight"] if s["freight"] is not None else "N/A — Missing Source Data", M2)
        b.f(ws, r, 22, f'=IF(ISNUMBER(U{r}),G{r}-U{r},{NA})',
            s["cost_excl"] if s["cost_excl"] is not None else "N/A — Missing Source Data", M2)
        b.put(ws, r, 23, s["local_transport"] if s["local_transport"] is not None else "N/A — Missing Source Data", M2)
        b.f(ws, r, 24, f'=IF(ISNUMBER(W{r}),H{r}-W{r},{NA})',
            s["local_excl"] if s["local_excl"] is not None else "N/A — Missing Source Data", M2)
        b.f(ws, r, 25, f'=IF(AND(ISNUMBER(V{r}),ISNUMBER(X{r})),X{r}-V{r},{NA})',
            s["saving_excl_mt"] if s["saving_excl_mt"] is not None else "N/A — Missing Source Data", M2)
        b.f(ws, r, 26, f'=IF(ISNUMBER(Y{r}),Y{r}*F{r},{NA})',
            s["saving_excl"] if s["saving_excl"] is not None else "N/A — Missing Source Data", M0)
        b.put(ws, r, 27, (f"Local Market Price — {s['local_age']} Days Old — STALE" if s["local_stale"]
                          else ("Same-day" if s["local_age"] == 0 else f"{s['local_age']} days old")),
              color=AMBER if s["local_stale"] else None, bold=s["local_stale"])
        b.put(ws, r, 28, s.get("storage_start") or "Not recorded")
        b.f(ws, r, 29, f'=IF(ISNUMBER(X{r}),X{r}*F{r},"")',
            (s["local_excl"] * s["mt"]) if s["local_excl"] is not None else None, M0)
        b.f(ws, r, 30, f'=IF(ISNUMBER(Y{r}),F{r},0)', s["mt"] if s["saving_excl_mt"] is not None else 0, M0)
        # Local equivalent premium = (local excl. transport − port costs) ÷ FX ÷ factor − CBOT that day.
        nv = lambda x: "N/A — Missing Source Data" if x is None else x
        has = s.get("premium_factor") is not None
        b.put(ws, r, 31, nv(s.get("day_cbot")) if has else "No CBOT for this commodity", M2)
        b.put(ws, r, 32, s.get("day_cbot_date") or "")
        b.put(ws, r, 33, nv(s.get("day_fx")), "0.0000"); b.put(ws, r, 34, s.get("day_fx_date") or "")
        b.put(ws, r, 35, nv(s.get("port_costs")) if has else "", M2)
        b.put(ws, r, 36, s.get("premium_factor") if has else "", "0.0000")
        b.f(ws, r, 37, f'=IF(AND(ISNUMBER(X{r}),ISNUMBER(AE{r}),ISNUMBER(AG{r}),ISNUMBER(AJ{r})),(X{r}-N(AI{r}))/AG{r}/AJ{r}-AE{r},{NA})',
            nv(s.get("eq_premium")), M2, bold=True)
        b.f(ws, r, 38, f'=IF(AND(ISNUMBER(X{r}),ISNUMBER(AE{r}),ISNUMBER(AG{r}),ISNUMBER(AJ{r})),X{r}/AG{r}/AJ{r}-AE{r},{NA})',
            nv(s.get("eq_premium_gross")), M2)
        b.put(ws, r, 39, s["contract_premium"] if s.get("contract_premium") is not None else ("Flat price" if has else ""), M2)
        b.f(ws, r, 40, f'=IF(AND(ISNUMBER(AK{r}),ISNUMBER(AM{r})),AK{r}-AM{r},{NA})', nv(s.get("premium_gap")), M2)
        b.put(ws, r, 41, s.get("premium_unit") or "")
    b.signs(ws, f"L{FIRST}:M{last}")
    b.signs(ws, f"Y{FIRST}:Z{last}")
    b.signs(ws, f"AN{FIRST}:AN{last}")
    orange = PatternFill("solid", fgColor="FFE3B3")
    ws.conditional_formatting.add(f"H{FIRST}:J{last}", FormulaRule(formula=[f"$J{FIRST}>5"], fill=orange))
    ws.conditional_formatting.add(f"AA{FIRST}:AA{last}", FormulaRule(formula=[f"$J{FIRST}>5"], fill=orange))
    return last


# ── 4b. Contract Savings Summary ─────────────────────────────────────────

def _contract_savings(b, m, last):
    ws = b.ws(5)
    rows = [c for c in m["contracts"] if c["consumed_mt"] > 0]
    rows.sort(key=lambda c: -(c["hist"]["saving"] or -1e18))
    b.title(ws, "Contract Savings Summary", "Active consumption performance: one row per contract with "
            "consumption (zero-consumption contracts stay in the Contract Master only). Every figure is a "
            "formula over Daily Consumption Allocation and Savings vs Local.")
    heads = ["Contract ID", "Contract Number", "Supplier", "Commodity", "Consumed MT", "Compared MT",
             "Not Compared MT", "Landed Cost Incl. Transport EGP/MT", "Wtd Avg Local Incl. Transport EGP/MT",
             "Total Savings Incl. Transport EGP", "Savings Incl. Transport EGP/MT", "Savings %", "First Consumption",
             "Last Consumption", "CIF Basis", "Result", "Final Assessment", "Base",
             "Total Savings Excl. Transport EGP", "Savings Excl. Transport EGP/MT", "Landed Cost Excl. Transport EGP/MT",
             "Wtd Avg Local Excl. Transport EGP/MT", "Storage Start Date", "Receipt Status", "Premium", "Premium Unit",
             "Futures Month", "CBOT Fixing", "Fixing Type", "Final CIF USD/MT", "Pricing / Fixing Date",
             "Pricing Status", "Oldest Local Quote Used (days)", "MT on Stale Local Quotes (>5 d)"]
    end = b.table(ws, "ContractSavings", heads, len(rows),
                  [9, 22, 16, 11, 10, 10, 10, 11, 11, 13, 10, 8, 12, 12, 13, 9, 90, 8,
                   13, 11, 12, 12, 12, 22, 9, 10, 10, 10, 14, 10, 14, 13, 10, 11])
    for i, c in enumerate(rows):
        r = FIRST + i
        h = c["hist"]
        b.put(ws, r, 1, c["id"]); b.put(ws, r, 2, c["contract"]); b.put(ws, r, 3, c.get("supplier"))
        b.put(ws, r, 4, c.get("commodity"))
        b.f(ws, r, 5, f"=SUMIFS({col(3, 'G', last['al'])},{col(3, 'B', last['al'])},A{r})", c["consumed_mt"], M0)
        b.f(ws, r, 6, f"=SUMIFS({col(4, 'F', last['sv'])},{col(4, 'B', last['sv'])},A{r})", h["priced_mt"], M0)
        b.f(ws, r, 7, f"=E{r}-F{r}", c["consumed_mt"] - h["priced_mt"], M0)
        b.put(ws, r, 8, blank(c.get("landed_egp_mt")), M2)
        b.f(ws, r, 9, f'=IF(F{r}=0,"",SUMIFS({col(4, "R", last["sv"])},{col(4, "B", last["sv"])},A{r})/F{r})',
            h["local_wavg"], M2)
        b.f(ws, r, 10, f'=IF(F{r}=0,"",SUMIFS({col(4, "M", last["sv"])},{col(4, "B", last["sv"])},A{r}))',
            h["saving"], M0, bold=True)
        b.f(ws, r, 11, f'=IF(F{r}=0,"",J{r}/F{r})', h["saving_mt"], M2)
        b.f(ws, r, 12, f'=IF(F{r}=0,"",J{r}/(I{r}*F{r}))',
            (h["saving_pct"] / 100) if h["saving_pct"] is not None else None, PCT)
        b.put(ws, r, 13, c["first_date"]); b.put(ws, r, 14, c["last_date"])
        b.put(ws, r, 15, "Final CIF" if c.get("cif_final") else "Provisional CIF")
        res = ("Not compared" if h["saving"] is None else "Saving" if h["saving"] >= 0 else "Loss")
        b.f(ws, r, 16, f'=IF(J{r}="","Not compared",IF(J{r}>=0,"Saving","Loss"))', res)
        b.put(ws, r, 17, c["assessment"], wrap=False)
        b.put(ws, r, 18, c["base"])
        SVc = lambda L_: col(4, L_, last["sv"])
        mtx = f'SUMIFS({SVc("AD")},{SVc("B")},A{r})'
        b.f(ws, r, 19, f'=IF({mtx}=0,{NA},SUMIFS({SVc("Z")},{SVc("B")},A{r}))',
            h["saving_excl"] if h["saving_excl"] is not None else "N/A — Missing Source Data", M0, bold=True)
        b.f(ws, r, 20, f'=IF(ISNUMBER(S{r}),S{r}/{mtx},{NA})',
            h["saving_excl_mt"] if h["saving_excl_mt"] is not None else "N/A — Missing Source Data", M2)
        b.put(ws, r, 21, h["cost_excl"] if h["cost_excl"] is not None else "N/A — Missing Source Data", M2)
        b.f(ws, r, 22, f'=IF({mtx}=0,{NA},SUMIFS({SVc("AC")},{SVc("B")},A{r})/{mtx})',
            h["local_excl_wavg"] if h["local_excl_wavg"] is not None else "N/A — Missing Source Data", M2)
        nr = c["storage_status"] == "Not recorded"
        b.put(ws, r, 23, c["storage_status"], color=AMBER if nr else None, bold=nr)
        b.put(ws, r, 24, c["receipt_status"])
        f_ = c["fix"]
        b.put(ws, r, 25, blank(c.get("premium")) if f_["eligible"] else "—  flat price", M2)
        b.put(ws, r, 26, f_["unit"] if f_["eligible"] else "")
        b.put(ws, r, 27, ((c.get("contract_month") or {}).get("code") or c.get("futures_month") or "") if f_["eligible"] else "")
        b.put(ws, r, 28, blank(f_["value"]) if f_["fixed"] else (f"Indicative {f_['indicative']:,.2f}" if f_.get("indicative") else ""), M2)
        b.put(ws, r, 29, ("Implied CBOT Fixing" if f_["type"] == "Implied" else
                          "Recorded fixing (lots)" if f_["type"] == "Recorded" else f_["type"]))
        b.put(ws, r, 30, blank(c.get("cif_usd_mt")) if c["pricing_class"] == "Fixed" else "", M2)
        b.put(ws, r, 31, f_["pricing_date"] or ("Not recorded" if f_["fixed"] else ""))
        b.put(ws, r, 32, c["pricing_class"])
        b.put(ws, r, 33, blank(h["max_local_age"]))
        b.f(ws, r, 34, f'=SUMIFS({SVc("F")},{SVc("B")},A{r},{SVc("J")},">5")', h["stale_mt"], M0)
    b.signs(ws, f"J{FIRST}:L{end}")
    b.signs(ws, f"S{FIRST}:T{end}")
    ws.conditional_formatting.add(f"AG{FIRST}:AH{end}", FormulaRule(formula=[f"$AG{FIRST}>5"],
                                                                    fill=PatternFill("solid", fgColor="FFE3B3")))
    m["_cs_rows"] = rows
    return end


# ── 5. CBOT Fixed vs Consumption ─────────────────────────────────────────

def _cbot(b, m):
    ws = b.ws(6)
    rows = m["cbot"]
    n = m["options"]["neutral_cbot"]
    b.title(ws, "CBOT Fixed vs Consumption", "CBOT Difference = consumption-day CBOT − contract CBOT fixed; "
            "USD/MT = difference × conversion factor; USD = USD/MT × consumed MT. Compared only on days the "
            "CBOT stored in the app for the consumption day itself (the futures month is only the month the "
            "premium is quoted over; the series' month that day is in the last column). A historical benchmark "
            "comparison, not a proven "
            "avoidable loss.")
    heads = ["Contract Number", "Contract ID", "Supplier", "Commodity", "Consumption Date", "Consumed MT",
             "CIF USD/MT", "Contract Premium", "Price Unit", "Conversion Factor", "CBOT Fixed", "Fixing Type",
             "Futures Month", "CBOT on Consumption Date", "Market Price Date", "Benchmark Status",
             "CBOT Difference", "Difference USD/MT", "Total Advantage / Missed USD", "FX",
             "Total Advantage / Missed EGP", "Result", "CBOT x MT", "Base", "Series Month That Day", "Month"]
    last = b.table(ws, "CbotVsConsumption", heads, len(rows),
                   [22, 9, 16, 11, 12, 9, 9, 9, 10, 9, 10, 11, 9, 11, 12, 22, 10, 10, 13, 8, 14, 16, 12, 8, 20, 9])
    for i, x in enumerate(rows):
        r = FIRST + i
        for j, v in enumerate([x["contract"], x["contract_id"], x["supplier"], x["commodity"], x["date"],
                               x["mt"], x["cif"], x["premium"], x["unit"], x["factor"], x["fix"],
                               x["fix_type"], x["futures"], x["cbot"], x["cbot_date"], x["benchmark"]], start=1):
            b.put(ws, r, j, blank(v), M0 if j == 6 else (M2 if j in (7, 8, 11, 14) else
                                                          ("0.0000" if j == 10 else None)))
        b.f(ws, r, 17, f'=IF(P{r}="{MATCHED}",N{r}-K{r},"")', x["diff"], M2)
        b.f(ws, r, 18, f'=IF(Q{r}="","",Q{r}*J{r})', x["usd_mt"], M2)
        b.f(ws, r, 19, f'=IF(R{r}="","",R{r}*F{r})', x["usd"], M0)
        b.put(ws, r, 20, blank(x["fx"]), "0.0000")
        b.f(ws, r, 21, f'=IF(OR(S{r}="",T{r}=""),"",S{r}*T{r})', x["egp"], M0)
        b.f(ws, r, 22, f'=IF(Q{r}="","Not assessable",IF(Q{r}>={n},"Advantage",IF(Q{r}<=-{n},"Missed opportunity","Neutral")))',
            x["result"])
        b.f(ws, r, 23, f'=IF(Q{r}="","",N{r}*F{r})', (x["cbot"] * x["mt"]) if x["assessable"] else None, M0)
        b.put(ws, r, 24, x["base"])
        b.put(ws, r, 25, x.get("series_month") or "")
        b.put(ws, r, 26, x["date"][:7])
    b.signs(ws, f"Q{FIRST}:U{last}")
    return last


# ── 6. Contract CBOT Performance ─────────────────────────────────────────

def _cbot_performance(b, m, last):
    ws = b.ws(7)
    n = m["options"]["neutral_cbot"]
    rows = [c for c in m["contracts"] if c["fix"]["fixed"] and c["cbot"]["consumed_mt"] > 0]
    rows.sort(key=lambda c: (c["cbot"]["usd"] is None, c["cbot"]["usd"] or 0))
    b.title(ws, "Contract CBOT Performance", "One row per fixed contract with consumption. Average CBOT is "
            "weighted by the tonnes consumed each day, against the CBOT stored for that day; the result is the "
            "sum of those days' differences — never the whole contract quantity. Not assessable where no CBOT is stored.")
    heads = ["Contract ID", "Contract", "Supplier", "Commodity", "MT Consumed", "MT Assessed", "Coverage %",
             "CIF $/MT", "Premium", "Price Unit", "CBOT Fixed", "Fixing Type", "Avg CBOT During Consumption",
             "Difference", "Advantage USD", "Missed Opportunity USD", "Net Result USD", "Net Result EGP",
             "Assessment", "Base"]
    end = b.table(ws, "ContractCbot", heads, len(rows),
                  [9, 22, 16, 11, 10, 10, 9, 9, 9, 10, 10, 11, 12, 10, 12, 13, 12, 13, 18, 8])
    C = lambda letter: col(6, letter, last["cb"])
    for i, c in enumerate(rows):
        r = FIRST + i
        cb, f = c["cbot"], c["fix"]
        b.put(ws, r, 1, c["id"]); b.put(ws, r, 2, c["contract"]); b.put(ws, r, 3, c.get("supplier"))
        b.put(ws, r, 4, c.get("commodity"))
        b.f(ws, r, 5, f"=SUMIFS({C('F')},{C('B')},A{r})", cb["consumed_mt"], M0)
        b.f(ws, r, 6, f'=SUMIFS({C("F")},{C("B")},A{r},{C("P")},"{MATCHED}")', cb["assessed_mt"], M0)
        b.f(ws, r, 7, f'=IF(E{r}=0,"",F{r}/E{r})', cb["coverage"], PCT)
        b.put(ws, r, 8, blank(c.get("cif_usd_mt")), M2); b.put(ws, r, 9, blank(c.get("premium")), M2)
        b.put(ws, r, 10, f["unit"]); b.put(ws, r, 11, f["value"], M2, bold=True); b.put(ws, r, 12, f["type"])
        b.f(ws, r, 13, f'=IF(F{r}=0,"",SUMIFS({C("W")},{C("B")},A{r})/F{r})', cb["avg_cbot"], M2)
        b.f(ws, r, 14, f'=IF(M{r}="","",M{r}-K{r})', cb["diff"], M2)
        b.f(ws, r, 15, f'=SUMIFS({C("S")},{C("B")},A{r},{C("S")},">0")', cb["advantage_usd"], M0)
        b.f(ws, r, 16, f'=SUMIFS({C("S")},{C("B")},A{r},{C("S")},"<0")', cb["missed_usd"], M0)
        b.f(ws, r, 17, f'=IF(F{r}=0,"",O{r}+P{r})', cb["usd"], M0, bold=True)
        b.f(ws, r, 18, f'=IF(F{r}=0,"",SUMIFS({C("U")},{C("B")},A{r}))', cb["egp"], M0)
        b.f(ws, r, 19, f'=IF(F{r}=0,"Not assessable",IF(N{r}>={n},"Advantage",IF(N{r}<=-{n},"Missed opportunity","Neutral")))',
            cb["assessment"])
        b.put(ws, r, 20, c["base"])
    b.signs(ws, f"N{FIRST}:R{end}")
    m["_cp_rows"] = rows
    return end


# ── 6b. Monthly Trend ────────────────────────────────────────────────

def _monthly(b, m, last):
    ws = b.ws(12)
    place(b, ws.title, SHEETS[7])
    b.title(ws, "Monthly Trend", "Savings vs local and the CBOT benchmark by consumption month — SUMIFS over "
            "Savings vs Local and CBOT Fixed vs Consumption (Month columns). The two are separate measures.")
    months = sorted({s["date"][:7] for s in m["savings"]} | {x["date"][:7] for x in m["cbot"] if x["assessable"]})
    heads = ["Month", "Savings vs Local Incl. Transport EGP", "Compared MT", "Savings Incl. EGP/MT", "CBOT Net USD",
             "CBOT Advantage USD", "CBOT Missed USD", "CBOT MT", "Savings vs Local Excl. Transport EGP"]
    end = b.table(ws, "MonthlyTrend", heads, len(months), [10, 16, 12, 12, 14, 15, 15, 11, 16])
    SV = lambda c_: col(4, c_, last["sv"])
    CB = lambda c_: col(6, c_, last["cb"])
    for i, mo in enumerate(months):
        r = FIRST + i
        sv = [s for s in m["savings"] if s["date"][:7] == mo]
        cb = [x for x in m["cbot"] if x["assessable"] and x["date"][:7] == mo]
        sav, mt = sum(s["saving"] for s in sv), sum(s["mt"] for s in sv)
        b.put(ws, r, 1, mo)
        b.f(ws, r, 2, f"=SUMIFS({SV('M')},{SV('T')},A{r})", sav, M0)
        b.f(ws, r, 3, f"=SUMIFS({SV('F')},{SV('T')},A{r})", mt, M0)
        b.f(ws, r, 4, f'=IF(C{r}=0,"",B{r}/C{r})', (sav / mt) if mt else None, M2)
        b.f(ws, r, 5, f"=SUMIFS({CB('S')},{CB('Z')},A{r})", sum(x["usd"] for x in cb), M0)
        b.f(ws, r, 6, f'=SUMIFS({CB("S")},{CB("Z")},A{r},{CB("S")},">0")', sum(x["usd"] for x in cb if x["usd"] > 0), M0)
        b.f(ws, r, 7, f'=SUMIFS({CB("S")},{CB("Z")},A{r},{CB("S")},"<0")', sum(x["usd"] for x in cb if x["usd"] < 0), M0)
        b.f(ws, r, 8, f"=SUMIFS({CB('F')},{CB('Z')},A{r})", sum(x["mt"] for x in cb), M0)
        b.f(ws, r, 9, f"=SUMIFS({SV('Z')},{SV('T')},A{r})", sum((x["saving_excl"] or 0) for x in sv), M0)
    b.signs(ws, f"B{FIRST}:I{end}")
    for title, col_i, anchor in (("Savings vs local by month, EGP", 2, "J4"), ("CBOT net result by month, USD", 5, "J22")):
        ch = BarChart()
        ch.title, ch.legend = title, None
        ch.y_axis.number_format = '#,##0.0,,"M"' if col_i == 2 else '#,##0,"k"'
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.add_data(Reference(ws, min_col=col_i, min_row=4, max_row=end), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=FIRST, max_row=end))
        ch.height, ch.width = 7.5, 16
        ws.add_chart(ch, anchor)


# ── 7. Open Contracts vs Local Today ─────────────────────────────────────

def open_bucket(o):
    """Which total a row counts in: confirmed, indicative / provisional, or neither."""
    if o["class"] == "Fixed price":
        return "Confirmed"
    if o["class"] == "Not assessable" or "incomplete" in o["class"]:
        return "Excluded — missing source data"
    return "Indicative / provisional"


def _open(b, m):
    ws = b.ws(8)
    rows = m["open"]
    b.title(ws, "Open Contracts vs Local Today", "Current open valuations: unconsumed tonnes (available stock + "
            "outstanding) at the latest local price, with and without inland transport. Fixed CIF is kept as "
            "recorded; unpriced contracts use the latest stored CBOT + premium (indicative); FX is Form 4 when "
            "recorded, else the latest daily FX (provisional). Contracts not yet received are a current market "
            "valuation. Potential, not realized — only 'Confirmed' rows count in the confirmed total.")
    heads = ["Contract Number", "Contract ID", "Supplier", "Commodity", "Origin", "Contract Status",
             "Contract Qty MT", "Delivered MT", "Consumed MT", "Available Stock MT", "Outstanding MT",
             "Remaining (Unconsumed) MT", "Storage Start Date", "Receipt Status", "Premium", "Premium Unit",
             "Futures Month", "Fixed CBOT", "Latest CBOT", "Latest CBOT Date", "CBOT Basis Note",
             "CIF USD/MT", "CIF Basis", "FX EGP/USD", "FX Basis", "FX Date", "Form 4 Status",
             "Discharge EGP/MT", "Clearance EGP/MT", "Import Inland Transport EGP/MT",
             "Imported Cost Excl. Transport EGP/MT", "Imported Cost Incl. Transport EGP/MT",
             "Local Price Excl. Transport EGP/MT", "Local Transport EGP/MT", "Local Price Incl. Transport EGP/MT",
             "Local Price Date", "Local Price Age (days)", "Local Price Age Status",
             "Savings Excl. Transport EGP/MT", "Savings Incl. Transport EGP/MT",
             "Total Potential Excl. Transport EGP", "Total Potential Incl. Transport EGP",
             "Pricing Status", "Valuation Class", "Open Contract", "Delivery Confirmed", "Base", "Issues",
             "Totals Bucket"]
    last = b.table(ws, "OpenVsLocal", heads, len(rows),
                   [20, 9, 16, 11, 11, 9, 10, 10, 10, 10, 10, 11, 12, 22, 9, 10, 9, 10, 10, 11, 44,
                    10, 34, 9, 30, 11, 14, 10, 10, 11, 12, 12, 12, 10, 12, 11, 9, 30, 12, 12, 14, 14,
                    12, 30, 8, 9, 8, 30, 22])
    na = "N/A — Missing Source Data"
    v = lambda x: na if x is None else x
    for i, o in enumerate(rows):
        r = FIRST + i
        for j, val in enumerate([o["contract"], o["contract_id"], o["supplier"], o["commodity"], o["origin"],
                                 o["status"], o["qty_mt"], o["delivered_mt"], o["consumed_mt"], o["available_mt"]],
                                start=1):
            b.put(ws, r, j, blank(val), M0 if j >= 7 else None)
        b.f(ws, r, 11, f"=G{r}-H{r}", o["outstanding_mt"], M0)
        b.f(ws, r, 12, f"=J{r}+K{r}", o["unconsumed_mt"], M0, bold=True)
        nr = o["storage_status"] == "Not recorded"
        b.put(ws, r, 13, o["storage_status"], color=AMBER if nr else None, bold=nr)
        b.put(ws, r, 14, o["receipt_status"], color=AMBER if not o["received"] else None)
        flat = o["pricing_class"] in ("Flat-priced", "Not CBOT-based")
        b.put(ws, r, 15, "— flat price" if flat else blank(o["premium"]), M2)
        b.put(ws, r, 16, "" if flat else o["unit"]); b.put(ws, r, 17, "" if flat else o["futures"])
        b.put(ws, r, 18, blank(o["fixed_cbot"]) if o["fixed_cbot"] is not None else ("Not fixed" if not flat else ""), M2)
        b.put(ws, r, 19, "" if flat else blank(o["latest_cbot"]), M2)
        b.put(ws, r, 20, "" if flat else o["latest_cbot_date"])
        b.put(ws, r, 21, "" if flat else o["cbot_note"], color=AMBER if "proxy" in (o["cbot_note"] or "") else None)
        b.put(ws, r, 22, v(o["cif"]), M2); b.put(ws, r, 23, o["cif_basis"])
        b.put(ws, r, 24, v(o["fx"]), "0.0000")
        b.put(ws, r, 25, o["fx_basis"], color=AMBER if o["fx_basis"].startswith("Latest") else None)
        b.put(ws, r, 26, o["fx_date"]); b.put(ws, r, 27, o["form4_status"])
        b.put(ws, r, 28, "Not recorded" if o.get("disc_missing") else o["discharge"], M2,
              color=AMBER if o.get("disc_missing") else None)
        b.put(ws, r, 29, "Not recorded" if o.get("clr_missing") else o["clearance"], M2,
              color=AMBER if o.get("clr_missing") else None)
        b.put(ws, r, 30, v(o["freight"]), M2)
        b.f(ws, r, 31, f'=IF(AND(ISNUMBER(V{r}),ISNUMBER(X{r})),V{r}*X{r}+N(AB{r})+N(AC{r}),{NA})', v(o["cost_excl"]), M2)
        b.f(ws, r, 32, f'=IF(AND(ISNUMBER(AE{r}),ISNUMBER(AD{r})),AE{r}+AD{r},{NA})', v(o["cost_incl"]), M2)
        b.put(ws, r, 33, v(o["local_excl"]), M2); b.put(ws, r, 34, v(o["local_transport"]), M2)
        b.f(ws, r, 35, f'=IF(AND(ISNUMBER(AG{r}),ISNUMBER(AH{r})),AG{r}+AH{r},{NA})', v(o["local_incl"]), M2)
        b.put(ws, r, 36, o["local_date"] or na)
        b.put(ws, r, 37, blank(o["local_age"]))
        b.put(ws, r, 38, (na if o["local_age"] is None else
                          f"Local Market Price — {o['local_age']} Days Old — STALE" if o["local_stale"] else
                          ("Same-day" if o["local_age"] == 0 else f"{o['local_age']} days old")),
              color=AMBER if o["local_stale"] or o["local_age"] is None else None, bold=o["local_stale"])
        b.f(ws, r, 39, f'=IF(AND(ISNUMBER(AG{r}),ISNUMBER(AE{r})),AG{r}-AE{r},{NA})', v(o["saving_excl_mt"]), M2)
        b.f(ws, r, 40, f'=IF(AND(ISNUMBER(AI{r}),ISNUMBER(AF{r})),AI{r}-AF{r},{NA})', v(o["saving_mt"]), M2)
        b.f(ws, r, 41, f'=IF(ISNUMBER(AM{r}),AM{r}*L{r},{NA})', v(o["potential_excl"]), M0, bold=True)
        b.f(ws, r, 42, f'=IF(ISNUMBER(AN{r}),AN{r}*L{r},{NA})', v(o["potential"]), M0, bold=True)
        b.put(ws, r, 43, o["pricing_class"]); b.put(ws, r, 44, o["class"])
        b.put(ws, r, 45, "Yes" if o["is_open"] else "No")
        b.put(ws, r, 46, "Yes" if o["received"] else "No")
        b.put(ws, r, 47, o["base"]); b.put(ws, r, 48, "; ".join(list(o["issues"]) + list(o.get("notes") or [])))
        b.put(ws, r, 49, open_bucket(o))
    b.signs(ws, f"AM{FIRST}:AP{last}")
    orange = PatternFill("solid", fgColor="FFE3B3")
    ws.conditional_formatting.add(f"AG{FIRST}:AL{last}", FormulaRule(formula=[f"$AK{FIRST}>5"], fill=orange))
    ws.conditional_formatting.add(f"AR{FIRST}:AR{last}", FormulaRule(formula=[f'$AW{FIRST}<>"Confirmed"'],
                                                                   fill=PatternFill("solid", fgColor="FFF4D6")))
    ws.freeze_panes = "C5"
    return last


# ── 8. Excluded & Missing Data ───────────────────────────────────────────

def _excluded(b, m):
    ws = b.ws(9)
    order = ["Unsourced consumption", "Cleared by stock count", "Logged entry not used", "No landed cost",
             "No usable local price", "Contract not fixed", "CBOT comparison not assessable",
             "Delivery not verified", "Duplicate record dropped", "Stock count vs FIFO model"]
    rows = sorted(m["excluded"], key=lambda e: (order.index(e["category"]) if e["category"] in order else 99,
                                                 e.get("date") or "", e.get("contract_id") or ""))
    b.title(ws, "Excluded & Missing Data", "Every record kept out of the financial results, and why. Nothing "
            "here is counted as zero savings or filled with an assumed value.")
    heads = ["Category", "Date", "Contract", "Contract ID", "Commodity", "Quantity MT", "Reason", "Source"]
    last = b.table(ws, "Excluded", heads, len(rows), [28, 12, 22, 10, 10, 10, 90, 22])
    for i, e in enumerate(rows):
        r = FIRST + i
        for j, v in enumerate([e["category"], e.get("date"), e.get("contract"), e.get("contract_id"),
                               e.get("commodity"), e.get("mt"), e["reason"], e.get("source")], start=1):
            b.put(ws, r, j, blank(v), M0 if j == 6 else None)
    return last


# ── 9. Validation & Sources ──────────────────────────────────────────────

def _validation(b, m, last):
    ws = b.ws(10)
    b.title(ws, "Validation & Sources", f"Checks run before export · model checksum {m['checksum']} (the HTML "
            "dashboard prints the same checksum when it renders the same figures).")
    ready = m["management_ready"]
    b.put(ws, 3, 1, "MANAGEMENT-READY — all checks passed" if ready else
          "NOT MANAGEMENT-READY — see the failed checks below", bold=True, color="FFFFFF",
          fill=GREEN if ready else RED)
    for j, h in enumerate(("Check", "Result", "Detail"), start=1):
        b.put(ws, 4, j, h, bold=True, color="FFFFFF", fill=NAVY)
    r = FIRST
    for v in m["validation"]:
        b.put(ws, r, 1, v["check"]); b.put(ws, r, 2, "PASS" if v["passed"] else "FAIL", bold=True,
                                           color=GREEN if v["passed"] else RED)
        b.put(ws, r, 3, v["detail"])
        r += 1
    r += 1
    b.put(ws, r, 1, "AUDIT — Total Consumed MT = Sourced Allocated MT + Excluded / Unallocated MT", bold=True, color=NAVY)
    r += 1
    for j, h in enumerate(("Commodity", "Total Consumed MT", "Sourced Allocated MT", "Excluded / Unallocated MT",
                           "Difference MT"), start=1):
        b.put(ws, r, j, h, bold=True, color="FFFFFF", fill=NAVY)
    r += 1
    for a in m["audit"]:
        b.put(ws, r, 1, a["base"]); b.put(ws, r, 2, a["consumed_mt"], M0)
        b.f(ws, r, 3, f'=SUMIFS({col(3, "G", last["al"])},{col(3, "M", last["al"])},A{r})'
                      f'+SUMIFS({col(11, "F", last["lq"])},{col(11, "L", last["lq"])},A{r})', a["allocated_mt"], M0)
        b.f(ws, r, 4, f'=SUMIFS({col(9, "F", last["ex"])},{col(9, "E", last["ex"])},A{r},{col(9, "A", last["ex"])},"Unsourced consumption")'
                      f'+SUMIFS({col(9, "F", last["ex"])},{col(9, "E", last["ex"])},A{r},{col(9, "A", last["ex"])},"Cleared by stock count")',
            a["excluded_mt"], M0)
        b.f(ws, r, 5, f"=B{r}-C{r}-D{r}", a["difference_mt"], M0)
        r += 1
    r += 1
    b.put(ws, r, 1, "SOURCES", bold=True, color=NAVY)
    r += 1
    src = m.get("sources") or {}
    lines = [("Data as of", m["as_of"]), ("Generated", m["generated"]), ("Consumption method", m["walk_mode"]),
             ("Daily rates (MT/day)", ", ".join(f"{k} {v:g}" for k, v in (m.get("rates") or {}).items())),
             ("CBOT history", src.get("cbot_source", "cbot_history"))]
    lines += [(f"CBOT {k}", f"{v[0]} → {v[1]} ({v[2]} days)") for k, v in (src.get("cbot") or {}).items()]
    lines += [(f"Latest local price {k}", v) for k, v in sorted((src.get("local_latest") or {}).items())]
    lines += [("FX history", " → ".join(src.get("fx") or [])),
              ("One-day CBOT spikes stepped over", ", ".join(f"{k} {v}" for k, v in (src.get("spikes") or {}).items()))]
    for k, v in lines:
        b.put(ws, r, 1, k); b.put(ws, r, 2, str(v)); r += 1
    r += 1
    b.put(ws, r, 1, "METHOD", bold=True, color=NAVY)
    r += 1
    o = m["options"]
    for t in [
        "Consumption: FIFO allocation of the configured daily rate from each delivery, oldest cargo first "
        "(FIFO Estimated). Logged withdrawals are used as Actual only when they are credible; see Excluded.",
        "Unsourced consumption (days with no cargo or local purchase in stock) is excluded, never allocated.",
        "Local price: the origin-matched local market quote in force on the day — the last quote entered stays "
        "in force until the next one, because no new quote means no change. Its date and age are shown on every "
        f"row; a quote more than {o.get('carry_days', 5)} days old is highlighted orange as "
        "'Local Market Price — N Days Old — STALE' (still used, flagged). No quote at all = missing, excluded.",
        "Two transport scenarios, each side with its own recorded transport: Including transport = landed cost "
        "incl. inland freight vs local quote + local transport; Excluding transport = landed cost excl. inland "
        "freight vs the local quote itself. Ocean freight stays inside CIF in both. Missing transport is shown "
        "as 'N/A — Missing Source Data', never as zero, and that row is left out of that scenario's total.",
        "Landed cost: CIF × FX + discharge + clearance (+ inland freight in the incl. scenario). FX = Form 4 once "
        "entered (all commodities except SBM); historical rows keep their recorded basis; open contracts without "
        "Form 4 use the latest daily market FX, labelled 'Latest market FX — provisional (Form 4 pending)', and "
        "switch automatically once Form 4 is entered. Discharge / clearance not recorded are flagged, not invented.",
        "Final CIF = (fixed CBOT + premium) × the commodity's factor. Open unpriced contracts are valued "
        "indicatively at (latest stored CBOT + premium) × factor, labelled a proxy when the latest stored series "
        "is not the contract's own futures month. Flat-priced contracts keep their contract CIF; no premium is "
        "invented for them.",
        "Storage start date is shown as recorded ('Not recorded' when absent) and never assumed from the delivery "
        "date. A contract with no storage start and no receipt is not treated as received: it is valued at the "
        "current market and kept out of the FIFO stock; any consumption already allocated to it is flagged for "
        "reconciliation, not deleted.",
        "Totals: historical savings (consumed tonnes), current open potential (unconsumed tonnes: confirmed = "
        "fixed price with Form 4; indicative / provisional kept separate; not assessable excluded) and the CBOT "
        "benchmark are three separate measures, never added together.",
        "CBOT fixed: recorded pricing lots (tonne-weighted), else implied = CIF ÷ factor − premium "
        "(CORN 0.3937; SOYBEAN, WHEAT 0.36745; SBM 1.1023, premium in $/short ton). Not shown as fixed for "
        "unpriced contracts. No price-adjustment fields are recorded in the database.",
        "Consumption-day CBOT: the CBOT stored in the app for each consumption day (the nearest futures "
        "month that day, recorded per row). The contract's futures month is the month its premium is quoted "
        "over and does not restrict the comparison. "
        f"Prints more than {o['cbot_max_age_days']} days old and one-day spikes are not used.",
        f"Result: Advantage when CBOT − fixed ≥ {o['neutral_cbot']:g} price unit, Missed opportunity when ≤ "
        f"−{o['neutral_cbot']:g}, else Neutral. A historical benchmark, not an executable alternative.",
        "Savings vs local and the CBOT result are separate measures and are never added together.",
    ]:
        b.put(ws, r, 1, t); r += 1
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 14


# ── 0. Management Summary ────────────────────────────────────────────────

def _summary(b, m, last):
    ws = b.ws(0)
    k, qd = m["kpis"], m["quality"]
    ready = m["management_ready"]
    b.title(ws, "Procurement Savings — Management Summary",
            f"Data as of {m['as_of']} · generated {m['generated']} · checksum {m['checksum']} · {k['contracts']} contracts")
    failed = [v["check"] for v in m["validation"] if not v["passed"]]
    b.put(ws, 3, 2, ("Management-ready: all validation checks passed" if ready else
                     "NOT management-ready — failed: " + "; ".join(failed)),
          bold=True, color="FFFFFF", fill=GREEN if ready else RED)
    ws.merge_cells("B3:I3")
    SV = lambda c: col(4, c, last["sv"])
    CB = lambda c: col(6, c, last["cb"])
    OP = lambda c: col(8, c, last["op"])
    cards = [
        ("Historical savings vs local — incl. transport, EGP", f"=SUM({SV('M')})", k["historical_saving_egp"],
         (f'="Excl. transport: EGP "&TEXT(SUM({SV("Z")}),"#,##0")&" · {k["historical_mt"]:,.0f} MT · actual, FIFO estimated"',
          f"Excl. transport: EGP {k['historical_saving_excl_egp']:,.0f} · {k['historical_mt']:,.0f} MT · actual, FIFO estimated")),
        ("Open potential vs local — confirmed, incl. transport, EGP", f'=SUMIFS({OP("AP")},{OP("AW")},"Confirmed")',
         k["open_fixed_potential_egp"],
         (f'="Excl. transport: EGP "&TEXT(SUMIFS({OP("AO")},{OP("AW")},"Confirmed"),"#,##0")&" · indicative / provisional EGP "&TEXT(SUMIFS({OP("AP")},{OP("AW")},"Indicative / provisional"),"#,##0")&" kept separate"',
          f"Excl. transport: EGP {k['open_fixed_potential_excl_egp']:,.0f} · indicative / provisional EGP "
          f"{k['open_indicative_potential_egp']:,.0f} kept separate")),
        ("Open contract outstanding, MT", f'=SUMIFS({OP("K")},{OP("AS")},"Yes",{OP("AT")},"Yes")',
         k["open_outstanding_mt"], f"Not yet delivered · {k['open_available_mt']:,.0f} MT delivered, in stock"),
        ("CBOT pricing advantage, USD", f'=SUMIFS({CB("S")},{CB("S")},">0")', k["cbot_advantage_usd"],
         "Benchmark — days the market was above our fixing"),
        ("CBOT missed opportunity, USD", f'=SUMIFS({CB("S")},{CB("S")},"<0")', k["cbot_missed_usd"],
         "Benchmark — days the market was below our fixing"),
        ("Net CBOT pricing difference, USD", f"=SUM({CB('S')})", k["cbot_net_usd"],
         f"EGP {k['cbot_net_egp']:,.0f} · {(k['cbot_coverage'] or 0) * 100:,.0f}% of fixed-contract tonnes with a stored CBOT"),
    ]
    for i, (label, formula, value, sub) in enumerate(cards):
        c0 = 2 + (i % 3) * 3
        r0 = 5 if i < 3 else 9
        b.put(ws, r0, c0, label, bold=True, color="FFFFFF", fill="1A6EBD", wrap=True)
        ws.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c0 + 1)
        b.f(ws, r0 + 1, c0, formula, value, M0)
        ws.cell(r0 + 1, c0).font = Font(bold=True, size=18, color=GREEN if (value or 0) >= 0 else RED)
        ws.merge_cells(start_row=r0 + 1, start_column=c0, end_row=r0 + 1, end_column=c0 + 1)
        if isinstance(sub, tuple):
            b.f(ws, r0 + 2, c0, sub[0], sub[1], italic=True, color=GREY, wrap=True)
        else:
            b.put(ws, r0 + 2, c0, sub, italic=True, color=GREY, wrap=True)
        ws.merge_cells(start_row=r0 + 2, start_column=c0, end_row=r0 + 2, end_column=c0 + 1)
        ws.row_dimensions[r0].height = 30
        ws.row_dimensions[r0 + 1].height = 28
        ws.row_dimensions[r0 + 2].height = 30
    for c_ in range(1, 14):
        ws.column_dimensions[get_column_letter(c_)].width = (3 if c_ in (1, 4, 7) else 21)

    r = 13
    b.put(ws, r, 2, "WHAT THIS SAYS", bold=True, color=NAVY)
    says = [
        f"Importing instead of buying locally {'saved' if k['historical_saving_egp'] >= 0 else 'cost'} "
        f"EGP {abs(k['historical_saving_egp']):,.0f} on {k['historical_mt']:,.0f} MT consumed "
        f"({k['historical_saving_egp_mt'] or 0:,.0f} EGP/MT), each tonne against the local price observed on its day.",
        f"Excluding transport on both sides the same tonnes saved EGP {k['historical_saving_excl_egp']:,.0f} "
        f"({k['historical_saving_excl_egp_mt'] or 0:,.0f} EGP/MT on {k['historical_mt_excl']:,.0f} MT).",
        f"Not-yet-consumed fixed-price tonnes would save EGP {k['open_fixed_potential_egp']:,.0f} incl. transport "
        f"(EGP {k['open_fixed_potential_excl_egp']:,.0f} excl.) at the latest local price; unpriced / provisional "
        f"contracts add an indicative EGP {k['open_indicative_potential_egp']:,.0f} incl. "
        f"(EGP {k['open_indicative_potential_excl_egp']:,.0f} excl.), kept separate.",
        f"On the consumption days the CBOT was above our fixing worth USD {k['cbot_advantage_usd']:,.0f} and below it "
        f"worth USD {abs(k['cbot_missed_usd']):,.0f} — net USD {k['cbot_net_usd']:,.0f}. A benchmark, not added to savings.",
        f"Excluded: {qd['excluded_unsourced_mt']:,.0f} MT of unsourced consumption, {qd['missing_cbot_mt']:,.0f} MT "
        "without a stored CBOT for the day — listed on Excluded & Missing Data. "
        f"{qd['stale_local']:,} day-rows ({qd['stale_local_mt']:,.0f} MT) used a local quote more than "
        f"{m['options'].get('carry_days', 5)} days old (orange); {qd['reconciliation_flags']} reconciliation flags.",
    ]
    for i, t in enumerate(says):
        b.put(ws, r + 1 + i, 2, f"{i + 1}.  {t}")

    r = 19
    b.put(ws, r, 2, "BY COMMODITY", bold=True, color=NAVY)
    heads = ["Commodity", "Historical savings incl. transport EGP", "Compared MT", "EGP/MT incl.",
             "Open potential confirmed incl. EGP", "Open indicative / provisional incl. EGP", "CBOT net USD",
             "CBOT assessed MT", "Historical savings excl. transport EGP", "EGP/MT excl.",
             "Open potential confirmed excl. EGP", "Open indicative / provisional excl. EGP"]
    for j, h in enumerate(heads):
        b.put(ws, r + 1, 2 + j, h, bold=True, color="FFFFFF", fill=NAVY, wrap=True)
    for i, c in enumerate(m["commodities"]):
        rr = r + 2 + i
        b.put(ws, rr, 2, c["base"], bold=True)
        b.f(ws, rr, 3, f"=SUMIFS({SV('M')},{SV('S')},B{rr})", c["saving_egp"], M0)
        b.f(ws, rr, 4, f"=SUMIFS({SV('F')},{SV('S')},B{rr})", c["priced_mt"], M0)
        b.f(ws, rr, 5, f'=IF(D{rr}=0,"",C{rr}/D{rr})', c["saving_egp_mt"], M0)
        b.f(ws, rr, 6, f'=SUMIFS({OP("AP")},{OP("AU")},B{rr},{OP("AW")},"Confirmed")', c["open_fixed_egp"], M0)
        b.f(ws, rr, 7, f'=SUMIFS({OP("AP")},{OP("AU")},B{rr},{OP("AW")},"Indicative / provisional")', c["open_indicative_egp"], M0)
        b.f(ws, rr, 10, f"=SUMIFS({SV('Z')},{SV('S')},B{rr})", c["saving_excl_egp"], M0)
        b.f(ws, rr, 11, f'=IF(SUMIFS({SV("AD")},{SV("S")},B{rr})=0,"",J{rr}/SUMIFS({SV("AD")},{SV("S")},B{rr}))',
            c["saving_excl_egp_mt"], M0)
        b.f(ws, rr, 12, f'=SUMIFS({OP("AO")},{OP("AU")},B{rr},{OP("AW")},"Confirmed")', c["open_fixed_excl_egp"], M0)
        b.f(ws, rr, 13, f'=SUMIFS({OP("AO")},{OP("AU")},B{rr},{OP("AW")},"Indicative / provisional")', c["open_indicative_excl_egp"], M0)
        b.f(ws, rr, 8, f"=SUMIFS({CB('S')},{CB('X')},B{rr})", c["cbot_net_usd"], M0)
        b.f(ws, rr, 9, f'=SUMIFS({CB("F")},{CB("X")},B{rr},{CB("P")},"{MATCHED}")', c["cbot_assessed_mt"], M0)
    com_last = r + 1 + len(m["commodities"])
    b.signs(ws, f"C{r + 2}:M{com_last}")

    r = com_last + 2
    b.put(ws, r, 2, "TOP 10 — LOCAL MARKET SAVINGS", bold=True, color=GREEN)
    b.put(ws, r, 6, "TOP 10 — CBOT MISSED OPPORTUNITY", bold=True, color=RED)
    for j, h in enumerate(("Contract", "Commodity", "Savings EGP")):
        b.put(ws, r + 1, 2 + j, h, bold=True)
    for j, h in enumerate(("Contract", "Commodity", "Net CBOT USD")):
        b.put(ws, r + 1, 6 + j, h, bold=True)
    top = [c for c in m["_cs_rows"] if (c["hist"]["saving"] or 0) > 0][:10]
    cs_index = {c["id"]: FIRST + i for i, c in enumerate(m["_cs_rows"])}
    for i, c in enumerate(top):
        rr = r + 2 + i
        b.put(ws, rr, 2, c["contract"]); b.put(ws, rr, 3, c.get("commodity"))
        b.f(ws, rr, 4, f"={q(SHEETS[5])}!J{cs_index[c['id']]}", c["hist"]["saving"], M0)
    cp_index = {c["id"]: FIRST + i for i, c in enumerate(m["_cp_rows"])}
    worst = [c for c in m["_cp_rows"] if (c["cbot"]["usd"] or 0) < 0][:10]
    for i, c in enumerate(worst):
        rr = r + 2 + i
        b.put(ws, rr, 6, c["contract"]); b.put(ws, rr, 7, c.get("commodity"))
        b.f(ws, rr, 8, f"={q(SHEETS[7])}!Q{cp_index[c['id']]}", c["cbot"]["usd"], M0)
    b.signs(ws, f"D{r + 2}:D{r + 12}")
    b.signs(ws, f"H{r + 2}:H{r + 12}")

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Historical savings vs local by commodity, EGP"
    chart.legend = None
    chart.y_axis.number_format = '#,##0.0,,"M"'
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.add_data(Reference(ws, min_col=3, min_row=20, max_row=com_last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=21, max_row=com_last))
    chart.height, chart.width = 7.5, 16
    ws.add_chart(chart, f"B{r + 14}")
